from io import BytesIO
import json
from pathlib import Path

import pytest
from flask import Flask, g
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

from core.daily_work import DailyWork
from core.sales_assistant import SalesAssistant, SalesAssistantError, sales_assistant_routes, MAX_ATTACHMENT
from core.store import dumps
from test_daily_work import store, user, enable_assignment


@pytest.fixture(autouse=True)
def stable_order_date(monkeypatch):
    # These fixtures explicitly use October 10/11 delivery dates. Keep their
    # creation date fixed so advancing wall time does not invalidate the data.
    from datetime import datetime
    class FixtureDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 4, 9, 0, tzinfo=tz)
    monkeypatch.setattr('core.sales_assistant.datetime', FixtureDatetime)


@pytest.fixture
def engine(store):
    daily=DailyWork(store,auto_project=False);enable_assignment(daily)
    with store.connect(write=True) as db:
        item={'profile_id':'P','warehouse_id':'08390','product_id':'product-1',
              'sku_id':'sku-1','base_unit':'1','unit_name':'','sku':'LS1000',
              'product_name':'螺栓','spec':'M20*80','model':'8.8级特殊热浸锌'}
        db.execute("INSERT INTO records(kind,id,sku,payload) VALUES ('inventory','record-1','LS1000',?)",(dumps(item),))
        db.execute("INSERT OR REPLACE INTO settings VALUES ('data_version',?)",(dumps('v1'),))
    return SalesAssistant(store)


def create(engine,actor=None,key='create'):
    actor=actor or user('sales',['sales'])
    return engine.create(actor,{'profile_id':'P','customer_name':'客户甲','idempotency_key':key})['session']


def edit_body(session,key='edit',**extra):
    return {'version':session['version'],'idempotency_key':key,'customer_order_no':'PO-123',
        'delivery_date':'2026-10-10','note':'销售确认','lines':[{
            'product_code':'LS1000','quantity':'3','unit_price':'1.25','unit':'支',
            'warehouse':'08390','delivery_date':'2026-10-11'}],**extra}


def _legacy_without_line_ids(engine, session, *, status=None, markers=None):
    """Represent a real older persisted row, including its frozen payload."""
    with engine.store.connect(write=True) as db:
        row=dict(db.execute('SELECT * FROM sales_assistant_sessions WHERE id=?',
                            (session['id'],)).fetchone())
        draft=json.loads(row['draft'])
        for line in draft['lines']: line.pop('id',None)
        row['draft']=dumps(draft)
        if row['confirmed_hash']:
            frozen=json.loads(row['confirmed_payload']);frozen['draft']=draft
            row['confirmed_payload']=dumps(frozen)
            from hashlib import sha256
            row['confirmed_hash']=sha256(row['confirmed_payload'].encode()).hexdigest()
        if status is not None: row['status']=status
        if markers: row.update(markers)
        db.execute('''UPDATE sales_assistant_sessions SET draft=?,status=?,confirmed_hash=?,
            confirmed_payload=?,submission_operation_id=?,erp_order_no=?,submitted_at=? WHERE id=?''',
            tuple(row[key] for key in ('draft','status','confirmed_hash','confirmed_payload',
                  'submission_operation_id','erp_order_no','submitted_at','id')))
        return row


def test_legacy_no_id_draft_reads_server_ids_and_can_edit_after_restart(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session,'before-upgrade'))['session']
    original=_legacy_without_line_ids(engine,draft)
    restarted=SalesAssistant(engine.store)
    viewed=restarted.detail(actor,session['id'])['session']
    assert viewed['version']==original['version']+1
    identity=viewed['draft']['lines'][0]['id']
    assert isinstance(identity,str) and len(identity)==32
    assert any(issue['code']=='legacy_line_identity_review' for issue in viewed['draft']['issues'])
    body={**viewed['draft'],'version':viewed['version'],'idempotency_key':'after-upgrade',
          'lines':[{**viewed['draft']['lines'][0],'quantity':'4','new_line':False}]}
    edited=restarted.edit(actor,session['id'],body)['session']
    assert edited['draft']['lines'][0]['id']==identity
    assert edited['draft']['lines'][0]['quantity']=='4' and not edited['draft']['issues']
    again=SalesAssistant(engine.store).detail(actor,session['id'])['session']
    assert again['version']==edited['version'] and again['draft']['lines'][0]['id']==identity
    with engine.store.connect() as db:
        assert db.execute("SELECT count(*) FROM sales_assistant_events WHERE action='migrate_line_identity'").fetchone()[0]==1


def test_legacy_duplicate_skus_get_distinct_stable_rows_without_guessing_client_ids(engine):
    actor=user('sales',['sales']);session=create(engine)
    rows=[{'product_code':'LS1000','quantity':str(qty),'unit_price':'1.13','unit':'支',
           'project_no':project,'summary_metadata_confirmed':True}
          for qty,project in ((1,'项目甲'),(2,'项目乙'))]
    draft=engine.edit(actor,session['id'],edit_body(session,'repeated-before',lines=rows))['session']
    legacy=_legacy_without_line_ids(engine,draft)
    restarted=SalesAssistant(engine.store);viewed=restarted.detail(actor,session['id'])['session']
    ids=[line['id'] for line in viewed['draft']['lines']]
    assert len(set(ids))==2
    assert [line['project_no'] for line in viewed['draft']['lines']]==['项目甲','项目乙']
    with pytest.raises(SalesAssistantError) as stale:
        restarted.edit(actor,session['id'],{**viewed['draft'],'version':legacy['version'],
            'idempotency_key':'stale-old-version'})
    assert stale.value.code=='sales_draft_version_changed'
    bad=[{**line,'id':'android-random-'+str(index),'new_line':False}
         for index,line in enumerate(viewed['draft']['lines'])]
    with pytest.raises(SalesAssistantError) as wrong:
        restarted.edit(actor,session['id'],{**viewed['draft'],'lines':bad,'version':viewed['version'],
            'idempotency_key':'random-not-guessed'})
    assert wrong.value.code=='sales_line_identity_changed'
    rows=[{**line,'quantity':str(index+3),'new_line':False}
          for index,line in enumerate(viewed['draft']['lines'])]
    edited=restarted.edit(actor,session['id'],{**viewed['draft'],'lines':rows,
        'version':viewed['version'],'idempotency_key':'stable-server-ids'})['session']
    assert [line['id'] for line in edited['draft']['lines']]==ids
    assert [line['project_no'] for line in edited['draft']['lines']]==['项目甲','项目乙']


def test_unsubmitted_legacy_confirmation_is_archived_and_requires_fresh_review(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session))['session']
    confirmed=engine.confirm(actor,session['id'],{'version':draft['version'],
        'idempotency_key':'old-confirmation'})['session']
    original=_legacy_without_line_ids(engine,confirmed)
    restarted=SalesAssistant(engine.store);viewed=restarted.detail(actor,session['id'])['session']
    assert viewed['status']=='draft' and viewed['version']==original['version']+1
    assert viewed['confirmed_version']==0 and viewed['confirmed_hash']=='' and not viewed['submit_available']
    with engine.store.connect() as db:
        archive=json.loads(db.execute("SELECT payload FROM sales_assistant_events WHERE action='migrate_line_identity'").fetchone()[0])['original_frozen']
        assert archive['confirmed_hash']==original['confirmed_hash']
        assert archive['confirmed_payload']==original['confirmed_payload']
        assert archive['draft']==original['draft']
    with pytest.raises(SalesAssistantError) as held:
        restarted.confirm(actor,session['id'],{'version':viewed['version'],
            'idempotency_key':'must-review-first'})
    assert held.value.code=='sales_draft_needs_review'
    reviewed=restarted.edit(actor,session['id'],{**viewed['draft'],'version':viewed['version'],
        'idempotency_key':'review-migrated'})['session']
    new=restarted.confirm(actor,session['id'],{'version':reviewed['version'],
        'idempotency_key':'new-confirmation'})['session']
    assert new['status']=='confirmed' and new['confirmed_hash']!=original['confirmed_hash']


def test_legacy_source_row_links_and_existing_ids_survive_migration(engine):
    actor=user('sales',['sales']);session=create(engine)
    source=engine.message(actor,session['id'],{'version':1,'idempotency_key':'legacy-source',
        'text':'LS1000,1,1.13\nLS1000,2,1.13'})['session']
    draft=engine.edit(actor,session['id'],{**source['draft'],'version':source['version'],
        'idempotency_key':'legacy-source-header','customer_order_no':'PO-123',
        'delivery_date':'2026-10-15'})['session']
    keep_id=draft['draft']['lines'][0]['id']
    legacy=_legacy_without_line_ids(engine,draft)
    with engine.store.connect(write=True) as db:
        body=json.loads(legacy['draft']);body['lines'][0]['id']=keep_id
        db.execute('UPDATE sales_assistant_sessions SET draft=? WHERE id=?',(dumps(body),session['id']))
    restarted=SalesAssistant(engine.store);viewed=restarted.detail(actor,session['id'])['session']
    lines=viewed['draft']['lines']
    assert lines[0]['id']==keep_id and lines[1]['id']!=keep_id
    by_source={line['source_ref']['id']:line['id'] for line in lines}
    assert all(report['line_id']==by_source[report['id']]
               for report in viewed['draft']['source_rows'] if report['status']!='ignored_with_reason')
    edited=restarted.edit(actor,session['id'],{**viewed['draft'],'version':viewed['version'],
        'idempotency_key':'linked-legacy-edit'})['session']
    assert [line['id'] for line in edited['draft']['lines']]==[line['id'] for line in lines]


def test_existing_queue_intent_prevents_legacy_id_migration_and_edit(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session))['session']
    original=_legacy_without_line_ids(engine,draft)
    engine.web_operations.enqueue(kind='sales_order_import',actor_id='sales',role='sales',
        profile_id='P',target_id='UL-legacy',key='legacy-queue-intent',
        payload={'session_id':session['id'],'frozen_session':{'id':session['id'],
                  'draft':original['draft']}})
    restarted=SalesAssistant(engine.store)
    with engine.store.connect() as db:
        assert dict(db.execute('SELECT * FROM sales_assistant_sessions WHERE id=?',
                               (session['id'],)).fetchone())==original
    with pytest.raises(SalesAssistantError) as error:
        restarted.edit(actor,session['id'],edit_body(draft,'legacy-queued-edit'))
    assert error.value.code=='sales_draft_locked'


@pytest.mark.parametrize('status',['submitting','submitted','submission_unknown','submission_blocked'])
def test_legacy_execution_states_never_rewrite_frozen_business_facts(engine,status):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session))['session']
    confirmed=engine.confirm(actor,session['id'],{'version':draft['version'],
        'idempotency_key':'confirmed-before-attempt'})['session']
    original=_legacy_without_line_ids(engine,confirmed,status=status)
    restarted=SalesAssistant(engine.store)
    with engine.store.connect() as db:
        current=dict(db.execute('SELECT * FROM sales_assistant_sessions WHERE id=?',
                                (session['id'],)).fetchone())
        assert current==original
        assert db.execute("SELECT count(*) FROM sales_assistant_events WHERE action='migrate_line_identity'").fetchone()[0]==0
    with pytest.raises(SalesAssistantError) as error:
        restarted.edit(actor,session['id'],edit_body(confirmed,'cannot-rewrite'))
    assert error.value.code=='sales_draft_locked'


@pytest.mark.parametrize('markers',[{'submission_operation_id':'existing-operation'},
                                  {'erp_order_no':'ERP-existing'},
                                  {'submitted_at':'2026-10-07T17:00:00+08:00'}])
def test_legacy_submit_references_protect_even_a_mislabeled_draft(engine,markers):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session))['session']
    original=_legacy_without_line_ids(engine,draft,markers=markers)
    restarted=SalesAssistant(engine.store)
    with engine.store.connect() as db:
        assert dict(db.execute('SELECT * FROM sales_assistant_sessions WHERE id=?',
                               (session['id'],)).fetchone())==original
    with pytest.raises(SalesAssistantError) as error:
        restarted.edit(actor,session['id'],edit_body(draft,'marker-protected'))
    assert error.value.code=='sales_draft_locked'


def file(raw,name='订单.pdf'):
    return FileStorage(stream=BytesIO(raw),filename=name,content_type='application/octet-stream')


def test_customer_choices_only_offer_confirmed_active_fixed_sales(engine):
    with engine.store.connect(write=True) as db:
        db.execute("INSERT INTO staff_users VALUES (?,?,?,?,?,?,?,?,?,?)",
                   ('sales2','sales2','销售乙','',dumps(['sales']),0,1,3,'',''))
        db.executemany("""INSERT INTO sales_customer_owners
            (profile_id,customer_name,owner_id,status,source,updated_at) VALUES (?,?,?,?,?,?)""",
            [('P','客户乙','sales2','assigned','customer_confirmed',''),
             ('P','待核对客户','sales','review','customer_confirmed',''),
             ('P','旧推断客户','sales','assigned','awaiting_confirmation',''),
             ('P','非销售客户','buyer','assigned','customer_confirmed',''),
             ('OTHER','其他账套客户','sales','assigned','customer_confirmed','')])
    choices=engine.customers(user('sales',['sales']),'P')
    assert choices['items']==[{'customer_name':'客户甲','owner_id':'sales','owner_name':'销售甲'}]
    assert choices['total']==1 and not choices['has_more']
    assert engine.customers(user('sales',['sales']),'P','乙')['items']==[]
    admin=engine.customers(user('admin',['sales'],admin=True),'P')
    assert {item['customer_name'] for item in admin['items']}=={'客户甲','客户乙'}
    with pytest.raises(SalesAssistantError): engine.customers(user('buyer',['purchase']),'P')


def test_review_confirm_unavailable_submit_never_changes_erp_and_keeps_frozen_version(engine):
    actor=user('sales',['sales']);session=create(engine)
    assert session['status']=='collecting' and not session['submit_available']
    request=edit_body(session)
    draft=engine.edit(actor,session['id'],request)['session']
    assert draft['draft']['customer_order_no']=='PO-123' and draft['draft']['delivery_date']=='2026-10-10'
    assert draft['draft']['lines'][0]['delivery_date']=='2026-10-11'
    assert draft['draft']['lines'][0]['resolved_identity']['sku_id']=='sku-1'
    assert draft['draft']['lines'][0]['unit']=='支'
    assert draft['draft']['lines'][0]['unit_evidence']=='user_confirmed_base_unit_1_zhi_2026_10_07'
    assert not draft['draft']['issues']
    confirmed=engine.confirm(actor,session['id'],{'version':draft['version'],'idempotency_key':'confirm'})['session']
    assert confirmed['status']=='confirmed' and confirmed['confirmed_version']==confirmed['version']
    assert len(confirmed['confirmed_hash'])==64
    with pytest.raises(SalesAssistantError) as error:
        engine.submit(actor,session['id'],{'version':confirmed['version'],'idempotency_key':'submit'})
    assert error.value.http_status==503 and error.value.code=='sales_submission_unavailable'
    assert engine.detail(actor,session['id'])['session']==confirmed
    with engine.store.connect() as db:
        assert db.execute("SELECT count(*) FROM records WHERE kind='orders'").fetchone()[0]==0
        assert db.execute('SELECT count(*) FROM jobs').fetchone()[0]==0
        frozen=json.loads(db.execute('SELECT confirmed_payload FROM sales_assistant_sessions').fetchone()[0])
        assert frozen['data_version']=='v1' and frozen['draft']==confirmed['draft']
    changed=engine.edit(actor,session['id'],edit_body(confirmed,'changed'))['session']
    assert changed['status']=='draft' and changed['confirmed_hash']==''


def test_idempotency_version_conflict_hold_and_atomic_rollback(engine):
    actor=user('sales',['sales']);session=create(engine)
    body={'version':1,'idempotency_key':'message','text':'LS1000,2,0.125'}
    first=engine.message(actor,session['id'],body)
    assert first['session']['draft']['lines'][0]['unit_price']=='0.125'
    assert engine.message(actor,session['id'],body)=={**first,'replayed':True}
    assert create(engine)==session
    with pytest.raises(SalesAssistantError) as error:
        engine.message(actor,session['id'],{**body,'text':'LS1000,3,0.125'})
    assert error.value.code=='idempotency_conflict'
    with pytest.raises(SalesAssistantError) as error:engine.edit(actor,session['id'],edit_body(session,'stale'))
    assert error.value.code=='sales_draft_version_changed'
    with engine.store.connect(write=True) as db:db.execute("INSERT OR REPLACE INTO settings VALUES ('restoration_hold','true')")
    assert engine.message(actor,session['id'],body)['replayed']
    with pytest.raises(SalesAssistantError) as error:
        engine.message(actor,session['id'],{'version':2,'idempotency_key':'hold','text':'LS1000,3,1'})
    assert error.value.code=='restoration_hold'
    with engine.store.connect() as db:
        assert db.execute('SELECT count(*) FROM sales_assistant_messages').fetchone()[0]==1
        assert db.execute('SELECT version FROM sales_assistant_sessions').fetchone()[0]==2


def test_customer_quote_defaults_inclusive_and_keeps_original_price_after_confirmation(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{
        'product_code':'LS1000','quantity':'3','unit_price':'1.13','unit':'支'}]))['session']
    line=draft['draft']['lines'][0]
    assert draft['draft']['tax_basis']=='tax_inclusive'
    assert (line['unit_price'],line['quoted_unit_price'])==('1.13','1.13')
    assert (line['tax_basis'],line['tax_rate'])==('tax_inclusive','13')
    assert (line['untaxed_unit_price'],line['untaxed_line_amount'])==('1.000','3.00')
    confirmed=engine.confirm(actor,session['id'],{
        'version':draft['version'],'idempotency_key':'quote-confirm'})['session']
    assert confirmed['draft']['lines'][0]==line


def test_explicit_untaxed_quote_rounds_unit_then_line_half_up(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{
        'product_code':'LS1000','quantity':'3','unit_price':'1.2345',
        'tax_basis':'未税','unit':'支'}]))['session']['draft']
    line=draft['lines'][0]
    assert not draft['issues']
    assert line['quoted_unit_price']=='1.2345'
    assert (line['tax_basis'],line['untaxed_unit_price'],line['untaxed_line_amount']) == (
        'tax_exclusive','1.235','3.71')


def test_unsupported_tax_rate_and_ambiguous_quote_basis_block_confirmation(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{
        'product_code':'LS1000','quantity':'3','unit_price':'1.13',
        'tax_basis':'不确定','tax_rate':'9','unit':'支'}]))['session']
    assert {'tax_basis_invalid','tax_rate_invalid'} <= {
        issue['code'] for issue in draft['draft']['issues']}
    with pytest.raises(SalesAssistantError) as error:
        engine.confirm(actor,session['id'],{
            'version':draft['version'],'idempotency_key':'tax-confirm'})
    assert error.value.code=='sales_draft_needs_review'


def test_legacy_draft_without_tax_fields_reconfirms_using_inclusive_default(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session))['session']
    with engine.store.connect(write=True) as db:
        old=json.loads(db.execute('SELECT draft FROM sales_assistant_sessions WHERE id=?',
                                  (session['id'],)).fetchone()[0])
        old.pop('tax_basis');old.pop('tax_rate')
        for line in old['lines']:
            for key in ('tax_basis','tax_rate','quoted_unit_price','untaxed_unit_price',
                        'untaxed_line_amount'):
                line.pop(key,None)
        db.execute('UPDATE sales_assistant_sessions SET draft=? WHERE id=?',
                   (dumps(old),session['id']))
    confirmed=engine.confirm(actor,session['id'],{
        'version':draft['version'],'idempotency_key':'legacy-confirm'})['session']
    assert confirmed['draft']['lines'][0]['tax_basis']=='tax_inclusive'


def test_current_fixed_owner_controls_all_reads_writes_and_customer_cannot_change(engine):
    actor=user('sales',['sales']);session=create(engine)
    with engine.store.connect(write=True) as db:db.execute("UPDATE staff_users SET roles='[\"sales\"]' WHERE id='buyer'")
    other=user('buyer',['sales'])
    assert engine.list(other)['total']==0
    with pytest.raises(SalesAssistantError) as error:engine.detail(other,session['id'])
    assert error.value.http_status==403
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],edit_body(session,customer_name='客户乙'))
    assert error.value.code=='customer_scope_immutable'
    with engine.store.connect(write=True) as db:db.execute("UPDATE sales_customer_owners SET owner_id='buyer' WHERE customer_name='客户甲'")
    with pytest.raises(SalesAssistantError):engine.detail(actor,session['id'])
    assert engine.detail(other,session['id'])['session']['fixed_sales_owner_id']=='buyer'
    assert engine.list(actor)['total']==0
    with engine.store.connect(write=True) as db:db.execute("UPDATE sales_customer_owners SET owner_id='' WHERE customer_name='客户甲'")
    admin=user('admin',['sales'],True)
    assert engine.detail(admin,session['id'])['session']['fixed_sales_owner_id']==''
    with pytest.raises(SalesAssistantError) as error:engine.edit(admin,session['id'],edit_body(session))
    assert error.value.code=='customer_owner_review'
    with pytest.raises(SalesAssistantError):engine.create(admin,{'profile_id':'P','customer_name':'新客户','idempotency_key':'new'})


def test_ambiguous_units_quantity_and_price_require_review_and_never_infer_identity(engine):
    actor=user('sales',['sales']);session=create(engine)
    with engine.store.connect(write=True) as db:
        item=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0]);item['warehouse_id']='other'
        db.execute("INSERT INTO records(kind,id,sku,payload) VALUES ('inventory','record-2','LS1000',?)",(dumps(item),))
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{'product_code':'LS1000','quantity':'-1','unit_price':''}]))['session']
    assert {issue['code'] for issue in draft['draft']['issues']}=={'quantity_invalid','unit_price_missing_or_invalid','product_identity_review'}
    with pytest.raises(SalesAssistantError) as error:engine.confirm(actor,session['id'],{'version':draft['version'],'idempotency_key':'bad-confirm'})
    assert error.value.code=='sales_draft_needs_review'
    resolved=engine.edit(actor,session['id'],edit_body(draft,'fixed'))['session']
    assert not resolved['draft']['issues']
    with engine.store.connect(write=True) as db:
        product=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0]);product['sku_id']='sku-new'
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(product),))
    with pytest.raises(SalesAssistantError) as error:engine.confirm(actor,session['id'],{'version':resolved['version'],'idempotency_key':'catalog-changed'})
    assert error.value.code=='sales_catalog_changed'


def test_all_base_unit_one_labels_are_pieces_and_raw_name_change_needs_reconfirmation(engine):
    actor=user('sales',['sales']);session=create(engine)
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{
        'product_code':'LS1000','quantity':'2','unit_price':'1.25','unit':'1'}]))['session']
    assert draft['draft']['lines'][0]['unit']=='支'
    assert draft['draft']['lines'][0]['unit_evidence']=='user_confirmed_base_unit_1_zhi_2026_10_07'
    assert not draft['draft']['issues']
    with engine.store.connect(write=True) as db:
        product=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        product['unit_name']='盒'
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(product),))
    with pytest.raises(SalesAssistantError) as error:
        engine.confirm(actor,session['id'],{'version':draft['version'],'idempotency_key':'stale-unit-confirm'})
    assert error.value.code=='sales_catalog_changed'
    changed=engine.edit(actor,session['id'],edit_body(draft,'unit-override',lines=[{
        'product_code':'LS1000','quantity':'2','unit_price':'1.25'}]))['session']
    line=changed['draft']['lines'][0]
    assert line['unit']=='支' and line['unit_evidence']=='user_confirmed_all_base_unit_1_zhi_2026_10_07'
    assert line['erp_unit_names']=={'unit_name':'盒','base_unit_name':'','unit':''}
    assert not changed['draft']['issues']
    assert engine.confirm(actor,session['id'],{'version':changed['version'],
        'idempotency_key':'unit-override-confirm'})['session']['status']=='confirmed'


def test_piece_rule_does_not_convert_another_erp_unit_id_or_typed_package_unit(engine):
    actor=user('sales',['sales']);session=create(engine)
    with engine.store.connect(write=True) as db:
        product=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        product.update(base_unit='2',unit_name='箱')
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(product),))
    draft=engine.edit(actor,session['id'],edit_body(session,'another-unit',lines=[{
        'product_code':'LS1000','quantity':'2','unit_price':'1.25','unit':'箱'}]))['session']['draft']
    assert {'unit_base_review','unit_mismatch'} <= {issue['code'] for issue in draft['issues']}


def test_customer_warehouse_name_resolves_to_exact_erp_warehouse_identity(engine):
    actor=user('sales',['sales']);session=create(engine)
    with engine.store.connect(write=True) as db:
        product=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        product['warehouse_name']='标准件仓'
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(product),))
    draft=engine.edit(actor,session['id'],edit_body(session,lines=[{
        'product_code':'LS1000','quantity':'2','unit_price':'1.25','warehouse':'标准件仓'}]))['session']['draft']
    assert not draft['issues']
    assert draft['lines'][0]['warehouse']=='08390'
    assert draft['lines'][0]['warehouse_name']=='标准件仓'


def test_attachment_blob_is_durable_verified_private_and_manual_edit_keeps_coverage_blocker(engine):
    actor=user('sales',['sales']);session=create(engine)
    request={'version':'1','idempotency_key':'upload'};raw=b'%PDF-1.7\nsource material'
    uploaded=engine.attachment(actor,session['id'],request,file(raw,'../../订单.pdf'))
    attachment=uploaded['session']['attachments'][0]
    assert attachment['name']=='订单.pdf' and attachment['status']=='needs_processing'
    assert len(attachment['sha256'])==64 and 'blob_name' not in attachment
    assert engine.attachment(actor,session['id'],request,file(raw,'../../订单.pdf'))['replayed']
    restarted=SalesAssistant(engine.store)
    assert restarted.attachment_file(actor,session['id'],attachment['id'])==(raw,'订单.pdf','application/pdf')
    blob=list(engine.blobs.glob('*.bin'))[0]
    assert blob.stat().st_mode & 0o777 == 0o600
    changed=engine.edit(actor,session['id'],edit_body(uploaded['session']))['session']
    assert changed['draft']['issues'] and changed['attachments'][0]['status']=='needs_processing'
    assert changed['draft']['source_rows'][0]['exclude_allowed'] is False
    with pytest.raises(SalesAssistantError) as error:
        engine.confirm(actor,session['id'],{'version':changed['version'],'idempotency_key':'confirm'})
    assert error.value.code=='sales_draft_needs_review'
    blob.write_bytes(b'damaged')
    with pytest.raises(SalesAssistantError) as error:engine.attachment_file(actor,session['id'],attachment['id'])
    assert error.value.code=='attachment_blob_unavailable'


def test_attachment_limits_and_wrong_mime_do_not_persist_orphan_data(engine):
    actor=user('sales',['sales']);session=create(engine);body={'version':'1','idempotency_key':'upload'}
    with pytest.raises(SalesAssistantError) as error:engine.attachment(actor,session['id'],body,file(b'x'*(MAX_ATTACHMENT+1)))
    assert error.value.http_status==413
    with pytest.raises(ValueError):engine.attachment(actor,session['id'],body,file(b'not a pdf'))
    assert list(engine.blobs.iterdir())==[]
    assert engine.detail(actor,session['id'])['session']['version']==1


def test_excel_requires_explicit_headers_and_holds_mismatched_customer_for_review(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active;sheet.append(['商品编码','数量','单价','单位','仓库','交期']);sheet.append(['LS1000',2,1.5,'支','08390','2026-10-15'])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{'version':'1','idempotency_key':'xlsx'},file(raw.getvalue(),'订单.xlsx'))['session']
    assert uploaded['attachments'][0]['status']=='parsed_for_review'
    assert uploaded['draft']['lines'][0]['delivery_date']=='2026-10-15'
    assert {item['code'] for item in uploaded['draft']['issues']}=={'customer_order_no_missing'}
    book=Workbook();sheet=book.active;sheet.append(['商品编码','数量','客户名称']);sheet.append(['LS1000',3,'客户乙'])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{'version':str(uploaded['version']),'idempotency_key':'other'},file(raw.getvalue(),'另单.xlsx'))['session']
    assert uploaded['attachments'][-1]['status']=='needs_review'
    assert len(uploaded['draft']['lines'])==1
    assert {item['code'] for item in uploaded['draft']['issues']}=={
        'customer_order_no_missing','source_row_needs_review'}
    assert uploaded['draft']['source_rows'][-1]['reason']=='source_customer_mismatch'


def test_excel_untaxed_price_header_is_explicit_and_conflicting_price_columns_are_held(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','数量','未税单价']);sheet.append(['LS1000',3,1.2345])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'untaxed-xlsx'},
        file(raw.getvalue(),'未税订单.xlsx'))['session']
    line=uploaded['draft']['lines'][0]
    assert uploaded['attachments'][0]['status']=='parsed_for_review'
    assert (line['tax_basis'],line['quoted_unit_price'],line['untaxed_unit_price'],
            line['untaxed_line_amount'])==('tax_exclusive','1.2345','1.235','3.71')
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],edit_body(uploaded,'old-client-edit'))
    assert error.value.code=='sales_tax_basis_required'
    assert engine.detail(actor,session['id'])['session']['version']==uploaded['version']
    conflicting=Workbook();sheet=conflicting.active
    sheet.append(['商品编码','数量','单价','未税单价']);sheet.append(['LS1000',1,1.13,1])
    raw=BytesIO();conflicting.save(raw)
    held=engine.attachment(actor,session['id'],{
        'version':str(uploaded['version']),'idempotency_key':'ambiguous-xlsx'},
        file(raw.getvalue(),'双单价.xlsx'))['session']
    assert held['attachments'][-1]['status']=='needs_review'
    assert len(held['draft']['lines'])==1


def test_binary_xls_with_explicit_headers_is_parsed_for_review(engine):
    actor=user('sales',['sales']);session=create(engine)
    raw=(Path(__file__).parent/'fixtures'/'sample_order.xls').read_bytes()
    uploaded=engine.attachment(actor,session['id'],
        {'version':'1','idempotency_key':'legacy-xls'},file(raw,'订单.xls'))['session']
    assert uploaded['attachments'][0]['status']=='parsed_for_review'
    assert uploaded['draft']['lines'][0]['product_code']=='LS1000'
    assert uploaded['draft']['lines'][0]['delivery_date']=='2026-10-15'


def test_text_rows_keep_bad_line_whitespace_format_and_stable_source_on_manual_match(engine):
    actor=user('sales',['sales']);session=create(engine)
    text='LS1000,2,1.13\n未给编码的六角螺栓 M20*80\nLS1000 3 1.13'
    uploaded=engine.message(actor,session['id'],{
        'version':1,'idempotency_key':'source-text','text':text})['session']
    draft=uploaded['draft']
    assert [item['row'] for item in draft['source_rows']]==[1,2,3]
    assert [item['product_code'] for item in draft['lines']]==['LS1000','','LS1000']
    assert len({item['source_ref']['id'] for item in draft['lines']})==3
    assert draft['source_rows'][1]['status']=='needs_review'
    assert draft['source_rows'][1]['reason']=='source_line_needs_match'
    assert 'product_identity_review' in draft['lines'][1]['issues']
    placeholder_id=draft['lines'][1]['id']
    body={**draft,'version':uploaded['version'],'idempotency_key':'source-review'}
    body['delivery_date']='2026-10-15'
    body['lines']=[dict(item) for item in draft['lines']]
    body['lines'][1].update(product_code='LS1000',quantity='1',unit_price='1.13',
                            replace_product_code=True)
    reviewed=engine.edit(actor,session['id'],body)['session']['draft']
    assert reviewed['lines'][1]['id']==placeholder_id
    assert reviewed['lines'][1]['source_ref']['match_basis']=='manual_code'
    assert reviewed['source_rows'][1]['line_id']==placeholder_id
    assert reviewed['source_rows'][1]['reason']=='manual_code_reviewed'
    assert reviewed['source_rows'][1]['status']=='parsed_for_review'
    assert reviewed['source_rows'][1]['preview']=='未给编码的六角螺栓 M20*80'


def test_text_source_keeps_original_physical_line_numbers(engine):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.message(actor,session['id'],{
        'version':1,'idempotency_key':'leading-blank',
        'text':'\nLS1000,2,1.13\n'})['session']
    assert uploaded['messages'][0]['text']=='\nLS1000,2,1.13\n'
    assert [(entry['row'],entry['reason']) for entry in uploaded['draft']['source_rows']]==[
        (1,'blank_row'),(2,'')]
    assert uploaded['draft']['lines'][0]['source_ref']['row']==2


def test_same_text_or_excel_content_with_new_key_does_not_duplicate_lines(engine):
    actor=user('sales',['sales']);session=create(engine)
    body={'version':1,'idempotency_key':'first-text','text':'LS1000,2,1.13'}
    first=engine.message(actor,session['id'],body)['session']
    again=engine.message(actor,session['id'],{
        **body,'version':first['version'],'idempotency_key':'second-text'})
    assert again['duplicate_of']==first['messages'][0]['id']
    assert again['session']['version']==first['version']
    assert len(again['session']['draft']['lines'])==1
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','数量','单价']);sheet.append(['LS1000',1,1.13])
    raw=BytesIO();book.save(raw)
    excel=engine.attachment(actor,session['id'],{
        'version':str(first['version']),'idempotency_key':'first-file'},
        file(raw.getvalue(),'订单.xlsx'))['session']
    duplicate=engine.attachment(actor,session['id'],{
        'version':str(excel['version']),'idempotency_key':'second-file'},
        file(raw.getvalue(),'另名订单.xlsx'))
    assert duplicate['duplicate_of']==excel['attachments'][0]['id']
    assert duplicate['session']['version']==excel['version']
    assert len(duplicate['session']['attachments'])==1
    assert len(duplicate['session']['draft']['lines'])==2


def test_dynamic_headers_and_multiple_sheets_keep_every_physical_row(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();one=book.active;one.title='客户报价'
    one.append(['客户报价单'])
    one.append(['请按行核对'])
    one.append(['物资名称','订购数量','ERP商品编号','含税单价','规格','单位'])
    one.append(['螺栓',2,'LS1000',1.13,'M20*80',''])
    one.append(['螺栓',3,'LS1000',1.13,'M20*80','支'])
    one.append(['合计',5])
    two=book.create_sheet('补单')
    two.append(['品名','订单数量','商品编码','销售单价'])
    two.append(['螺栓',4,'',1.13])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'multi-sheet'},file(raw.getvalue(),'报价.xlsx'))['session']
    draft=uploaded['draft']
    assert uploaded['attachments'][0]['status']=='parsed_for_review'
    assert len(draft['source_rows'])==8
    assert [(item['sheet'],item['row']) for item in draft['source_rows']]==[
        ('客户报价',1),('客户报价',2),('客户报价',3),('客户报价',4),
        ('客户报价',5),('客户报价',6),('补单',1),('补单',2)]
    assert [item['product_code'] for item in draft['lines']]==['LS1000','LS1000','']
    assert len({item['id'] for item in draft['lines']})==3
    assert len({item['source_ref']['id'] for item in draft['lines']})==3
    assert draft['source_rows'][0]['reason']=='before_header'
    assert draft['source_rows'][5]['reason']=='footer'
    assert draft['source_rows'][-1]['status']=='needs_review'
    assert draft['source_rows'][3]['unit_note'].startswith('来源单位空白；商品基本单位 1')
    assert draft['lines'][0]['unit']=='支'
    old_client={**draft,'version':uploaded['version'],'idempotency_key':'no-source-ids'}
    old_client['lines']=[{key:value for key,value in item.items() if key!='id'}
                         for item in draft['lines']]
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],old_client)
    assert error.value.code=='sales_line_identity_ambiguous'


def test_bad_customer_row_in_same_sheet_does_not_drop_other_rows(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','数量','单价','客户名称'])
    sheet.append(['LS1000',1,1.13,'客户甲'])
    sheet.append(['LS1000',2,1.13,'其他客户'])
    sheet.append(['LS1000',3,1.13,'客户甲'])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'mixed-customers'},
        file(raw.getvalue(),'混合客户.xlsx'))['session']
    assert uploaded['attachments'][0]['status']=='needs_review'
    assert [line['quantity'] for line in uploaded['draft']['lines']]==['1','3']
    assert [item['row'] for item in uploaded['draft']['source_rows']]==[1,2,3,4]
    assert uploaded['draft']['source_rows'][2]['reason']=='source_customer_mismatch'
    assert uploaded['draft']['source_rows'][2]['blocking'] is True
    assert 'source_row_needs_review' in {item['code'] for item in uploaded['draft']['issues']}


def test_footer_word_with_explicit_sku_is_still_a_source_line(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品名称','数量','商品编码','单价'])
    sheet.append(['合计',1,'LS1000',1.13])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'footer-with-code'},
        file(raw.getvalue(),'订单.xlsx'))['session']
    assert uploaded['draft']['lines'][0]['product_code']=='LS1000'
    assert uploaded['draft']['source_rows'][1]['line_id']==uploaded['draft']['lines'][0]['id']


def test_source_locator_survives_restart_edit_and_frozen_confirmation(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active;sheet.title='客户明细'
    sheet.append(['商品编码','数量','含税单价'])
    sheet.append(['LS1000',2,1.13])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'source-upload'},
        file(raw.getvalue(),'客户订单.xlsx'))['session']
    restarted=SalesAssistant(engine.store)
    viewed=restarted.detail(actor,session['id'])['session']
    source=viewed['draft']['lines'][0]['source_ref']
    assert source['sheet']=='客户明细' and source['row']==2
    body={**viewed['draft'],'version':viewed['version'],'idempotency_key':'source-edit'}
    body.update(customer_order_no='PO-123',delivery_date='2026-10-15')
    edited=restarted.edit(actor,session['id'],body)['session']
    assert edited['draft']['lines'][0]['source_ref']==source
    frozen=restarted.confirm(actor,session['id'],{
        'version':edited['version'],'idempotency_key':'source-confirm'})['session']
    assert frozen['status']=='confirmed'
    with engine.store.connect() as db:
        confirmed=json.loads(db.execute(
            'SELECT confirmed_payload FROM sales_assistant_sessions WHERE id=?',
            (session['id'],)).fetchone()[0])['draft']
    assert confirmed['lines'][0]['source_ref']==source
    assert confirmed['source_rows'][1]['line_id']==confirmed['lines'][0]['id']
    assert confirmed['source_rows'][1]['status']=='parsed_for_review'


def test_explicit_exclusion_of_orphan_text_row_is_audited_and_idempotent(engine):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.message(actor,session['id'],{
        'version':1,'idempotency_key':'orphan-message',
        'text':'LS1000,2,1.13\n备注：此行不是商品'})['session']
    draft=uploaded['draft'];orphan=draft['source_rows'][1]
    body={**draft,'version':uploaded['version'],'idempotency_key':'exclude-orphan',
          'customer_order_no':'PO-123','delivery_date':'2026-10-15',
          'lines':[draft['lines'][0]],
          'source_row_actions':[{'source_row_id':orphan['id'],'action':'exclude',
                                 'reason':'员工核对原文后确认是备注'}]}
    edited=engine.edit(actor,session['id'],body)['session']
    excluded=edited['draft']['source_rows'][1]
    assert excluded['status']=='excluded_by_user' and excluded['blocking'] is False
    assert excluded['original_reason']=='source_line_needs_match'
    assert excluded['original_preview']=='备注：此行不是商品'
    assert excluded['preview']==orphan['preview']
    assert excluded['exclusion_reason']=='员工核对原文后确认是备注'
    assert excluded['excluded_by']=='sales' and excluded['excluded_at']
    assert excluded['line_id']==draft['lines'][1]['id']
    assert engine.edit(actor,session['id'],body)['replayed'] is True
    assert engine.detail(actor,session['id'])['session']['version']==edited['version']
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],{**edited['draft'],'version':edited['version'],
            'idempotency_key':'exclude-orphan-again',
            'source_row_actions':body['source_row_actions']})
    assert error.value.code=='sales_source_row_action_invalid'
    confirmed=engine.confirm(actor,session['id'],{
        'version':edited['version'],'idempotency_key':'confirm-exclusion'})['session']
    assert confirmed['status']=='confirmed'
    assert confirmed['draft']['source_rows'][1]==excluded


def test_mixed_customer_row_can_be_explicitly_excluded_without_dropping_good_row(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','数量','含税单价','客户名称'])
    sheet.append(['LS1000',1,1.13,'客户甲'])
    sheet.append(['LS1000',2,1.13,'其他客户'])
    raw=BytesIO();book.save(raw)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'customer-file'},
        file(raw.getvalue(),'混合客户.xlsx'))['session']
    bad=uploaded['draft']['source_rows'][2]
    assert bad['line_id']=='' and bad['reason']=='source_customer_mismatch'
    body={**uploaded['draft'],'version':uploaded['version'],'idempotency_key':'customer-exclude',
          'customer_order_no':'PO-123','delivery_date':'2026-10-15',
          'source_row_actions':[{'source_row_id':bad['id'],'action':'exclude',
                                 'reason':'员工核对原件，属于其他客户'}]}
    edited=engine.edit(actor,session['id'],body)['session']
    assert len(edited['draft']['lines'])==1
    assert edited['draft']['source_rows'][2]['status']=='excluded_by_user'
    assert edited['draft']['source_rows'][2]['original_reason']=='source_customer_mismatch'
    assert not edited['draft']['issues']


def test_linked_source_row_requires_same_edit_line_removal_for_exclusion(engine):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.message(actor,session['id'],{
        'version':1,'idempotency_key':'two-lines',
        'text':'LS1000,1,1.13\nLS1000,2,1.13'})['session']
    draft=uploaded['draft'];second=draft['source_rows'][1]
    base={**draft,'version':uploaded['version'],'customer_order_no':'PO-123',
          'delivery_date':'2026-10-15'}
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],{**base,'idempotency_key':'remove-no-action',
            'lines':[draft['lines'][0]]})
    assert error.value.code=='sales_source_line_removed'
    action=[{'source_row_id':second['id'],'action':'exclude','reason':'重复来源行，人工核对排除'}]
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],{**base,'idempotency_key':'exclude-still-linked',
            'source_row_actions':action})
    assert error.value.code=='sales_source_row_still_linked'
    edited=engine.edit(actor,session['id'],{**base,'idempotency_key':'exclude-linked',
        'lines':[draft['lines'][0]],'source_row_actions':action})['session']
    assert len(edited['draft']['lines'])==1
    assert edited['draft']['source_rows'][1]['status']=='excluded_by_user'
    offered=dict(draft['lines'][1],id='new-local-id',new_line=True)
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],{**edited['draft'],'version':edited['version'],
            'idempotency_key':'reuse-excluded',
            'lines':[edited['draft']['lines'][0],offered]})
    assert error.value.code=='sales_source_row_excluded'


@pytest.mark.parametrize('actions', [
    [{'source_row_id':'forged','action':'exclude','reason':'人工核对'}],
    [{'source_row_id':'SOURCE','action':'exclude','reason':'人工核对'},
     {'source_row_id':'SOURCE','action':'exclude','reason':'再次排除'}],
    [{'source_row_id':'SOURCE','action':'exclude','reason':' '}],
    [{'source_row_id':'SOURCE','action':'exclude','reason':'x'*201}],
    [{'source_row_id':'SOURCE','action':'restore','reason':'人工核对'}],
])
def test_source_row_exclusion_rejects_forgery_duplicate_or_invalid_reason(engine, actions):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.message(actor,session['id'],{
        'version':1,'idempotency_key':'action-source',
        'text':'LS1000,1,1.13\n备注：非商品'})['session']
    draft=uploaded['draft'];source_id=draft['source_rows'][1]['id']
    actions=[{**entry,'source_row_id':source_id if entry['source_row_id']=='SOURCE'
              else entry['source_row_id']} for entry in actions]
    with pytest.raises((SalesAssistantError,ValueError)):
        engine.edit(actor,session['id'],{**draft,'version':uploaded['version'],
            'idempotency_key':'invalid-action','lines':[draft['lines'][0]],
            'source_row_actions':actions})
    assert engine.detail(actor,session['id'])['session']['version']==uploaded['version']


def test_unprocessed_pdf_has_no_excludable_source_row(engine):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.attachment(actor,session['id'],{
        'version':'1','idempotency_key':'pdf-to-review'},
        file(b'%PDF-1.7\nsource material','订单.pdf'))['session']
    assert uploaded['attachments'][0]['status']=='needs_processing'
    source=uploaded['draft']['source_rows'][0]
    assert source['scope']=='attachment_coverage' and source['exclude_allowed'] is False
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],{**uploaded['draft'],'version':uploaded['version'],
            'idempotency_key':'fake-pdf-source',
            'source_row_actions':[{'source_row_id':source['id'],
                                   'action':'exclude','reason':'声称已识别'}]})
    assert error.value.code=='sales_source_row_action_invalid'


def test_failed_excel_is_blocked_across_edit_restart_and_requires_explicit_reason(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','数量','单价','未税单价'])
    sheet.append(['LS1000',2,1.13,1.00])
    output=BytesIO();book.save(output)
    uploaded=engine.attachment(actor,session['id'],{'version':1,'idempotency_key':'bad-header'},
                               file(output.getvalue(),'错误双价格.xlsx'))['session']
    coverage=uploaded['draft']['source_rows'][0]
    assert coverage['scope']=='attachment_coverage' and coverage['blocking']
    restarted=SalesAssistant(engine.store)
    edited=restarted.edit(actor,session['id'],edit_body(uploaded,'manual-lines'))['session']
    with pytest.raises(SalesAssistantError) as error:
        restarted.confirm(actor,session['id'],{'version':edited['version'],'idempotency_key':'bad-confirm'})
    assert error.value.code=='sales_draft_needs_review'
    excluded=restarted.edit(actor,session['id'],{**edited['draft'],'version':edited['version'],
        'idempotency_key':'explicit-superseded-file','source_row_actions':[
            {'source_row_id':coverage['id'],'action':'exclude','reason':'原文件双价格列无效，已按客户修订稿手工重录并核对'}]})['session']
    assert not excluded['draft']['issues']
    saved=excluded['draft']['source_rows'][0]
    assert saved['excluded_by']=='sales' and saved['exclusion_reason'].startswith('原文件双价格')
    confirmed=restarted.confirm(actor,session['id'],{'version':excluded['version'],
                                                    'idempotency_key':'reviewed-confirm'})['session']
    assert confirmed['draft']['source_rows'][0]==saved


def test_source_product_difference_requires_bound_human_confirmation(engine):
    actor=user('sales',['sales']);session=create(engine)
    book=Workbook();sheet=book.active
    sheet.append(['商品编码','商品名称','规格','等级','表面处理','数量','单价'])
    sheet.append(['LS1000','垫圈','M8×30','4.8','发黑',1,1.13])
    output=BytesIO();book.save(output)
    uploaded=engine.attachment(actor,session['id'],{'version':1,'idempotency_key':'source-difference'},
                               file(output.getvalue(),'核对商品.xlsx'))['session']
    line=uploaded['draft']['lines'][0]
    assert line['source_ref']['original_grade']=='4.8'
    assert line['source_ref']['original_surface']=='发黑'
    assert line['source_review']['required'] and not line['source_review']['confirmed']
    request={**uploaded['draft'],'version':uploaded['version'],'idempotency_key':'fill-po',
             'customer_order_no':'REVIEW-PO','delivery_date':'2026-10-15'}
    edited=engine.edit(actor,session['id'],request)['session']
    with pytest.raises(SalesAssistantError):
        engine.confirm(actor,session['id'],{'version':edited['version'],'idempotency_key':'pre-review-confirm'})
    line=edited['draft']['lines'][0]
    line={**line,'source_match_confirmed':True,
          'source_match_evidence_hash':line['source_review']['evidence_hash'],
          'source_match_confirmation':{'confirmed_by':'forged'}}
    reviewed=engine.edit(actor,session['id'],{**edited['draft'],'lines':[line],
        'version':edited['version'],'idempotency_key':'human-source-review'})['session']
    reviewed_line=reviewed['draft']['lines'][0]
    assert not reviewed['draft']['issues']
    fact=reviewed_line['source_match_confirmation']
    assert fact['confirmed_by']=='sales' and fact['confirmed_at']
    confirmed=engine.confirm(actor,session['id'],{'version':reviewed['version'],
        'idempotency_key':'source-reviewed-confirm'})['session']
    assert confirmed['draft']['lines'][0]['source_match_confirmation']==fact
    with engine.store.connect(write=True) as db:
        payload=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        payload['spec']='M24*80'
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(payload),))
    changed=engine.edit(actor,session['id'],{**confirmed['draft'],'version':confirmed['version'],
        'idempotency_key':'after-catalog-change'})['session']
    assert not changed['draft']['lines'][0]['source_review']['confirmed']
    assert any(issue['code']=='source_product_match_unconfirmed' for issue in changed['draft']['issues'])


def test_pdf_worker_result_is_separate_adoption_atomic_and_duplicate_safe(engine):
    from test_sales_document_extract import _pdf
    actor=user('sales',['sales']);session=create(engine)
    raw=_pdf(b'BT /F1 12 Tf 20 100 Td (LS1000,1,1.13) Tj ET',
             b'BT /F1 12 Tf 20 100 Td (LS1000,2,1.13) Tj ET')
    uploaded=engine.attachment(actor,session['id'],{'version':1,'idempotency_key':'real-two-page-pdf'},
                               file(raw,'两页订单.pdf'))['session']
    attachment=uploaded['attachments'][0]
    assert attachment['document_processing']['status']=='queued'
    result=engine.documents.run_one('worker-token-'+'a'*32)
    assert result['status']=='ready' and result['candidate_count']==2 and result['page_count']==2
    # Extraction never moves a draft version or silently appends candidates.
    unchanged=engine.detail(actor,session['id'])['session']
    assert unchanged['version']==uploaded['version'] and unchanged['draft']['lines']==[]
    body={'version':unchanged['version'],'data_version':'v1','result_hash':result['result_hash'],
          'idempotency_key':'adopt-two-page'}
    adopted=engine.adopt_document(actor,session['id'],attachment['id'],body)['session']
    assert adopted['version']==unchanged['version']+1
    assert [line['quantity'] for line in adopted['draft']['lines']]==['1','2']
    assert [line['source_ref']['page'] for line in adopted['draft']['lines']]==[1,2]
    assert len({line['id'] for line in adopted['draft']['lines']})==2
    assert not any(row['scope']=='attachment_coverage' for row in adopted['draft']['source_rows'])
    assert adopted['attachments'][0]['document_processing']['status']=='adopted'
    assert engine.adopt_document(actor,session['id'],attachment['id'],body)['replayed']
    repeated=engine.adopt_document(actor,session['id'],attachment['id'],
        {**body,'version':adopted['version'],'idempotency_key':'adopt-again'})['session']
    assert repeated['version']==adopted['version'] and len(repeated['draft']['lines'])==2


def test_unreadable_pdf_and_legacy_attachment_do_not_lose_coverage(engine):
    actor=user('sales',['sales']);session=create(engine)
    uploaded=engine.attachment(actor,session['id'],{'version':1,'idempotency_key':'broken-pdf'},
                               file(b'%PDF-1.7\ninvalid','无效订单.pdf'))['session']
    attachment=uploaded['attachments'][0]
    result=engine.documents.run_one('worker-token-'+'b'*32)
    assert result['status']=='failed'
    with pytest.raises(SalesAssistantError) as error:
        engine.adopt_document(actor,session['id'],attachment['id'],
            {'version':uploaded['version'],'data_version':'v1','result_hash':result['result_hash'],
             'idempotency_key':'bad-adopt'})
    assert error.value.code=='sales_document_not_ready'
    with engine.store.connect(write=True) as db:
        old={**uploaded['draft'],'source_rows':[],'issues':[]}
        db.execute('UPDATE sales_assistant_sessions SET draft=? WHERE id=?',(dumps(old),session['id']))
        db.execute('DELETE FROM sales_document_jobs WHERE attachment_id=?',(attachment['id'],))
    restarted=SalesAssistant(engine.store)
    detail=restarted.detail(actor,session['id'])['session']
    assert detail['attachments'][0]['document_processing']['status']=='queued'
    changed=restarted.edit(actor,session['id'],edit_body(detail,'legacy-manual-edit'))['session']
    assert changed['draft']['source_rows'][0]['exclude_allowed'] is False
    assert changed['draft']['issues']


def test_explicit_new_local_line_and_product_replacement_recheck_identity(engine):
    actor=user('sales',['sales']);session=create(engine)
    with engine.store.connect(write=True) as db:
        record=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        record.update(sku='LS2000',product_id='product-2',sku_id='sku-2',
                      product_name='另一商品',spec='M30*40')
        db.execute("INSERT INTO records(kind,id,sku,payload) VALUES ('inventory','record-2','LS2000',?)",
                   (dumps(record),))
    body=edit_body(session)
    body['lines'][0].update(id='client-local-uuid',new_line=True,
                            project_no='旧项目',summary_metadata_confirmed=True)
    created=engine.edit(actor,session['id'],body)['session']
    old=created['draft']['lines'][0]
    assert old['id']!='client-local-uuid'
    assert old['manual_summary']=={'project_no':'旧项目'}
    with pytest.raises(SalesAssistantError) as error:
        engine.edit(actor,session['id'],edit_body(created,key='unknown-id',lines=[{
            **body['lines'][0],'id':'unknown','new_line':False}]))
    assert error.value.code=='sales_line_identity_changed'
    replacement=edit_body(created,key='replace-product')
    replacement['lines'][0].update(id=old['id'],product_code='LS2000',replace_product_code=True)
    changed=engine.edit(actor,session['id'],replacement)['session']['draft']['lines'][0]
    assert changed['id']==old['id']
    assert changed['resolved_identity']['sku_id']=='sku-2'
    assert changed['product_name']=='另一商品'
    assert changed['catalog_summary']['spec']=='M30*40'
    assert changed['manual_summary']=={}


def test_human_dimension_provenance_survives_quantity_edits_and_catalog_change_blocks(engine):
    owner=user('sales',['sales'])
    admin=user('admin',['sales'],True)
    session=create(engine,owner)
    first_body=edit_body(session,key='confirm-size')
    first_body['lines'][0].update(summary_dimension='M20X90',
                                 summary_metadata_confirmed=True)
    first=engine.edit(owner,session['id'],first_body)['session']
    first_line=first['draft']['lines'][0]
    assert not first['draft']['issues']
    assert first_line['summary_dimension_source']=='human_confirmed'
    provenance=(first_line['summary_metadata_actor'],first_line['summary_metadata_at'])
    assert provenance[0]=='sales' and provenance[1]

    def quantity_edit(current,actor,key,quantity):
        previous=current['draft']['lines'][0]
        body=edit_body(current,key=key)
        body['lines'][0].update(id=previous['id'],quantity=quantity,
            summary_dimension=previous['summary_dimension'],
            summary_dimension_source=previous['summary_dimension_source'],
            summary_metadata_confirmed=True,
            summary_metadata_actor='spoofed-client-actor',
            summary_metadata_at='2099-01-01T00:00:00')
        return engine.edit(actor,session['id'],body)['session']

    same=quantity_edit(first,owner,'same-actor-quantity','4')
    other=quantity_edit(same,admin,'other-actor-quantity','5')
    for result in (same,other):
        line=result['draft']['lines'][0]
        assert line['summary_dimension']=='M20X90'
        assert line['summary_dimension_source']=='human_confirmed'
        assert (line['summary_metadata_actor'],line['summary_metadata_at'])==provenance
        assert not result['draft']['issues']

    with engine.store.connect(write=True) as db:
        product=json.loads(db.execute("SELECT payload FROM records WHERE id='record-1'").fetchone()[0])
        product['spec']='M20*85'
        db.execute("UPDATE records SET payload=? WHERE id='record-1'",(dumps(product),))
    stale=quantity_edit(other,admin,'changed-catalog-quantity','6')
    stale_line=stale['draft']['lines'][0]
    assert 'summary_catalog_changed' in {issue['code'] for issue in stale['draft']['issues']}
    assert (stale_line['summary_metadata_actor'],stale_line['summary_metadata_at'])==provenance
    with pytest.raises(SalesAssistantError) as error:
        engine.confirm(admin,session['id'],
            {'version':stale['version'],'idempotency_key':'confirm-changed-catalog'})
    assert error.value.code=='sales_draft_needs_review'


def test_route_contract_and_http_submit_boundary(engine):
    app=Flask(__name__);app.testing=True;app.config['MAX_CONTENT_LENGTH']=12*1024*1024
    @app.before_request
    def authenticate():g.current_user=user('sales',['sales'])
    app.register_blueprint(sales_assistant_routes(engine));client=app.test_client();base='/api/v2/sales-assistant/sessions'
    response=client.post(base,json={'profile_id':'P','customer_name':'客户甲','idempotency_key':'create'})
    session=response.get_json()['session'];identity=session['id']
    assert response.status_code==200 and client.get(base).get_json()['total']==1
    response=client.put(base+'/'+identity+'/draft',json=edit_body(session));draft=response.get_json()['session']
    response=client.post(base+'/'+identity+'/confirm',json={'version':draft['version'],'idempotency_key':'confirm'})
    confirmed=response.get_json()['session']
    response=client.post(base+'/'+identity+'/submit',json={'version':confirmed['version'],'idempotency_key':'submit'})
    assert response.status_code==503 and response.get_json()['error']=='sales_submission_unavailable'
    assert client.get(base+'/'+identity).get_json()['session']['status']=='confirmed'
    response=client.post(base+'/'+identity+'/attachments',data={'version':str(confirmed['version']),
        'idempotency_key':'pdf','file':(BytesIO(b'%PDF-1.7\nabc'),'订单.pdf')})
    attachment=response.get_json()['session']['attachments'][0]
    assert client.get(base+'/'+identity+'/attachments/'+attachment['id']).data==b'%PDF-1.7\nabc'


def test_low_storage_refuses_new_attachment_without_dropping_session(engine,monkeypatch):
    actor=user('sales',['sales']);session=create(engine)
    monkeypatch.setattr('core.storage_capacity.capacity',lambda p:{'status':'critical'})
    with pytest.raises(SalesAssistantError) as err:
        engine.attachment(actor,session['id'],{'version':session['version'],'idempotency_key':'low-space'},file(b'%PDF-1.4\n%%EOF'))
    assert err.value.code=='storage_capacity_low'
    assert engine.detail(actor,session['id'])['session']['version']==session['version']
    with engine.store.connect() as db:
        assert db.execute('SELECT count(*) FROM sales_assistant_attachments').fetchone()[0]==0
