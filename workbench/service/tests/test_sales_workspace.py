import json
from types import SimpleNamespace
import pytest
from werkzeug.security import generate_password_hash
from app import create_app
from core.store import dumps
from core.sales_workspace import BASE
from core import sales_api_jobs as sales
from tools.sales_api_worker import run_one
from test_sales_assistant import engine,stable_order_date,create,edit_body
from test_daily_work import store,user


@pytest.fixture
def app(engine):
    with engine.store.connect(write=True) as db:
        db.execute("UPDATE staff_users SET password_hash=? WHERE id='admin'",(generate_password_hash('test-password'),))
        db.execute("UPDATE staff_users SET password_hash=? WHERE id='sales'",(generate_password_hash('staff-password'),))
    return create_app({'TESTING':True,'AUTO_IMPORT':False,'SOURCE_MODE':'api','PUBLIC_MODE':True,
        'ADMIN_PASSWORD_HASH':generate_password_hash('test-password'),'ERP_CLIENT':SimpleNamespace(last_auth_error=None),
        'API_SOURCE':object(),'SALES_WEB_ORIGIN':'http://localhost'},runtime_dir=engine.store.runtime)


def login(client):
    result=client.post(BASE+'/login',json={'username':'admin','password':'test-password'},headers={'Origin':'http://localhost'})
    assert result.status_code==200
    return {'Origin':'http://localhost','X-CSRF-Token':result.json['csrf']}


def test_owner_cookie_csrf_and_native_isolation(app):
    client=app.test_client()
    assert client.get('/sales').status_code==200
    assert client.get(BASE+'/bootstrap').status_code==401
    assert client.post(BASE+'/login',json={'password':'test-password'},headers={'Origin':'https://evil.invalid'}).status_code==403
    headers=login(client)
    assert 'HttpOnly' in client.get_cookie('ul_sales_session').http_only.__str__() or client.get_cookie('ul_sales_session').http_only
    assert client.get(BASE+'/bootstrap').json['handlers']==['李瑄','刘锋']
    assert client.get('/api/v2/tasks').status_code==401 # Cookie is not a general API credential.
    body={'profile_id':'P','customer_name':'客户甲','handler_name':'刘锋','idempotency_key':'one'}
    assert client.post('/api/v2/sales-assistant/sessions',json=body).status_code==403
    result=client.post('/api/v2/sales-assistant/sessions',json=body,headers=headers)
    assert result.status_code==200 and result.json['session']['draft']['handler_name']=='刘锋'
    assert result.json['session']['fixed_sales_owner_id']=='sales'
    assert client.post(BASE+'/logout',json={},headers=headers).status_code==200
    assert client.get(BASE+'/bootstrap').status_code==401


def test_staff_login_cannot_open_owner_workspace(app):
    client=app.test_client()
    result=client.post(BASE+'/login',json={'username':'sales','password':'staff-password'},headers={'Origin':'http://localhost'})
    assert result.status_code==403 and result.json['error']=='owner_only'


def confirmed(engine,lines=2):
    with engine.store.connect(write=True) as db:
        db.execute("UPDATE records SET payload=json_set(payload,'$.warehouse_id','29871293290508390','$.product_id','2851960081570977703','$.sku_id','7992537581237892200') WHERE kind='inventory'")
        db.execute("INSERT OR REPLACE INTO settings VALUES ('sales_api_policy',?)",(dumps({'enabled':True,'contract_accepted':True,'profile_id':'P','owner_user_id':'admin'}),))
    actor=user('admin',[],True);s=create(engine,actor)
    s=engine.edit(actor,s['id'],edit_body(s,handler_name='李瑄',delivery_date='2026-12-20',lines=[{
        'product_code':'LS1000','quantity':str(q),'unit_price':'1.25','unit':'支'} for q in range(3,3+lines)]))['session']
    assert not s['draft']['issues']
    s=engine.confirm(actor,s['id'],{'version':s['version'],'idempotency_key':'confirm'})['session']
    return actor,s


def binding(_client,model):
    return {'customer_id':'7991974216891337259','handler_id':'26775042030594223','products':{'LS1000':{
        'ptypeid':'2851960081570977703','fullname':'螺栓','sku':{'skuid':'7992537581237892200','prop1':'0','prop2':'0'}}}}


def test_exact_model_repeat_sku_and_shared_export(engine):
    actor,s=confirmed(engine)
    with engine.store.connect() as db:
        row=db.execute('SELECT * FROM sales_assistant_sessions WHERE id=?',(s['id'],)).fetchone();model=sales.normalized(db,row)
    docs=model['documents'];assert len(docs)==2 and docs[0]['number']!=docs[1]['number']
    assert docs[0]['unit_price']=='1.106' and docs[0]['amount']=='3.32' and docs[0]['gross_amount']=='3.75'
    assert docs[0]['summary'].endswith('#交期=2026-12-20')
    wire=sales.wire_document(model,docs[0],binding(None,model));encoded=sales.decimal_json(wire)
    assert 'comment' not in wire # User requested no additional note.
    assert '"ptypeid":2851960081570977703' in encoded and '"assdpprice":1.106' in encoded
    from openpyxl import load_workbook
    sheet=load_workbook(sales.workbook(model)).active
    assert sheet.cell(2,9).value=='1.106' and sheet.cell(2,14).value==wire['summary']
    assert sheet.max_column==14


def test_duplicate_submit_and_unknown_never_resend(engine):
    actor,s=confirmed(engine)
    body={'version':s['version'],'channel':'api','idempotency_key':'send'}
    first=engine.submit(actor,s['id'],body);again=engine.submit(actor,s['id'],body)
    assert first['operation_id']==again['operation_id']
    calls=[]
    def send(client,doc):calls.append(doc);return {'status':'unknown','message':'timeout'}
    run_one(engine.store,None,prepare=binding,send=send)
    run_one(engine.store,None,prepare=binding,send=send)
    assert len(calls)==1
    v=engine.detail(actor,s['id'])['session'];assert v['status']=='submission_unknown'
    assert [d['status'] for d in v['api_submission']['documents']]==['unknown','held']


def test_partial_acceptance_survives_restart(engine):
    actor,s=confirmed(engine)
    engine.submit(actor,s['id'],{'version':s['version'],'channel':'api','idempotency_key':'send'})
    calls=[]
    def send(client,doc):
        calls.append(doc)
        if len(calls)==2:raise RuntimeError('process lost')
        return {'status':'accepted','vchcode':'123'}
    with pytest.raises(RuntimeError):run_one(engine.store,None,prepare=binding,send=send)
    run_one(engine.store,None,prepare=binding,send=send)
    assert len(calls)==2
    v=engine.detail(actor,s['id'])['session'];assert v['api_submission']['documents'][0]['status']=='accepted'
    assert v['api_submission']['documents'][1]['status']=='unknown'


def test_api_not_enabled_preserves_confirmation(engine):
    actor,s=confirmed(engine)
    with engine.store.connect(write=True) as db:db.execute("DELETE FROM settings WHERE key='sales_api_policy'")
    with pytest.raises(Exception) as error:engine.submit(actor,s['id'],{'version':s['version'],'channel':'api','idempotency_key':'send'})
    assert error.value.code=='sales_api_not_accepted'
    assert engine.detail(actor,s['id'])['session']['status']=='confirmed'


def test_rule_versions_cannot_publish_without_test(app):
    client=app.test_client();headers=login(client)
    r=client.post(BASE+'/rules',headers=headers,json={'idempotency_key':'rule1','content':'不猜编码','problem':'螺栓','expected':'{"product_codes":[],"needs_clarification":true}'})
    assert r.status_code==200
    identity=r.json['id']
    r=client.post(BASE+'/rules/'+identity+'/publish',headers=headers,json={'idempotency_key':'pub','expected_active_id':''})
    assert r.status_code==400
    assert not client.get(BASE+'/rules').json['active_id']


def test_stopping_job_preserves_sources_and_does_not_reclaim(app,engine):
    from test_sales_codex_pipeline import enable
    enable(engine);actor=user('admin',[],True);s=create(engine,actor)
    result=engine.request_codex(actor,s['id'],{'version':s['version'],'text':'原始订单保留','idempotency_key':'ai'})
    s=result['session'];client=app.test_client();headers=login(client)
    body={'version':s['version'],'idempotency_key':'stop'}
    one=client.post(BASE+'/sessions/'+s['id']+'/stop',json=body,headers=headers)
    assert one.status_code==200 and one.json['session']['codex_job']['status']=='cancelled'
    assert one.json['session']['messages'][-1]['text']=='原始订单保留'
    assert client.post(BASE+'/sessions/'+s['id']+'/stop',json=body,headers=headers).json['replayed']
    assert engine.codex.claim('a'*32) is None


def test_rule_supplements_preserve_prior_version_and_rollback(app,engine):
    client=app.test_client();headers=login(client)
    def candidate(key,content):
        return client.post(BASE+'/rules',headers=headers,json={'idempotency_key':key,'content':content,
            'problem':'螺栓','expected':'{"product_codes":[],"needs_clarification":true}'}).json['id']
    first=candidate('a','保留客户原文')
    with engine.store.connect(write=True) as db:
        row=db.execute('SELECT content_hash FROM sales_rule_versions WHERE id=?',(first,)).fetchone()
        db.execute("UPDATE sales_rule_versions SET status='passed',evaluation=? WHERE id=?",(dumps({'candidate_hash':row[0],'base_hash':''}),first))
    r=client.post(BASE+'/rules/'+first+'/publish',headers=headers,json={'idempotency_key':'p','expected_active_id':''})
    assert r.status_code==200
    second=candidate('b','缺少单价时提问')
    versions=client.get(BASE+'/rules').json['versions']
    assert next(v for v in versions if v['id']==second)['content']=='保留客户原文\n\n缺少单价时提问'
    assert client.post(BASE+'/rules/rollback',headers=headers,json={'idempotency_key':'r','expected_active_id':first}).json['active_id']==''
