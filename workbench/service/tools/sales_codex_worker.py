"""One durable, owner-operated Codex sales draft job. No ERP or payment calls.

The root host prepares credential-free files, invokes a no-tools inference
child as ulcodex, retrieves current catalog/NAS evidence between phases, and
stores a proposal. Adopting and confirming that proposal are separate App
actions with the existing version/ownership checks.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import pwd
import re
import shutil
import signal
import stat
import subprocess
import sys
import uuid

from core.daily_work import _setting
from core.sales_codex_pipeline import (CodexJobError, SalesCodexPipeline,
                                       source_manifest, validate_output)
from core.sales_document_pipeline import _bounded_extract, _extraction_from_json
from core.sales_history_reference import SalesHistoryReference
from core.sales_skill_matching import build_matching_context, load_skill_rules
from core.sales_source_parser import parse_excel
from core.sales_dependencies import product_dependencies
from core.store import Store, dumps
from tools import sales_codex_runner as runner


SERVICE = Path(__file__).resolve().parents[1]
EXCEL_TYPES = {"application/vnd.ms-excel",
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp"}
FORMAL_HISTORY = "/var/lib/ul-sales-history-reference/20261007-formal/sales-matching.sqlite3"
FORMAL_HISTORY_SHA = "f59081526af28b5f31e73806e7b26f2a28f36e60d2aa4e38a646d8a892602e4a"
MAX_PACKET = 2 * 1024 * 1024
MAX_RENDERED_IMAGES = 10


class WorkerBlocked(ValueError):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class WorkerConfig:
    jobs_root: Path
    rules_root: Path
    codex_binary: Path = Path("/home/ulcodex/.local/bin/codex")
    codex_user: str | None = "ulcodex"
    timeout_seconds: int = 180
    document_python: Path | None = None
    pdftoppm_path: Path | None = None
    cancel_check: object = None


@contextmanager
def _database_identity(uid, gid):
    """Keep newly created SQLite WAL/SHM owned by the existing service user."""
    original_uid, original_gid = os.geteuid(), os.getegid()
    changed = original_uid == 0 and uid != original_uid
    try:
        if changed:
            os.setegid(gid)
            os.seteuid(uid)
        yield
    finally:
        if changed:
            os.seteuid(original_uid)
            os.setegid(original_gid)


class _ServiceIdentityStore:
    def __init__(self, store):
        self.runtime, self.path = store.runtime, store.path
        self._store = store
        info = self.path.stat()
        self._uid, self._gid = info.st_uid, info.st_gid

    @contextmanager
    def connect(self, write=False):
        with _database_identity(self._uid, self._gid), self._store.connect(write=write) as db:
            yield db


def _secure_bytes(path, *, maximum=MAX_PACKET, owner=None):
    path = Path(path)
    try:
        if path.resolve() != path:
            raise WorkerBlocked("sales_codex_file_path_invalid")
        descriptor = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
        with os.fdopen(descriptor, "rb") as stream:
            info = os.fstat(stream.fileno())
            if (not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_mode & 0o022
                    or (owner is not None and info.st_uid not in (0, owner))):
                raise WorkerBlocked("sales_codex_file_unsafe")
            raw = stream.read(maximum + 1)
            after = os.fstat(stream.fileno())
            if len(raw) > maximum:
                raise WorkerBlocked("sales_codex_file_limit")
            if (info.st_size, info.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                raise WorkerBlocked("sales_codex_file_changed")
            return raw
    except OSError as error:
        raise WorkerBlocked("sales_codex_file_unavailable") from error


def _identity(config):
    if config.codex_user is None:
        return os.getuid(), os.getgid(), None
    try:
        user = pwd.getpwnam(config.codex_user)
    except KeyError as error:
        raise WorkerBlocked("sales_codex_owner_account_unavailable") from error
    if os.getuid() != 0:
        raise WorkerBlocked("sales_codex_root_worker_required")
    return user.pw_uid, user.pw_gid, user


def _directory(path, uid, gid, *, traverse=False):
    path = Path(path)
    if not path.is_absolute() or path.resolve() != path:
        raise WorkerBlocked("sales_codex_job_path_invalid")
    path.mkdir(parents=True, exist_ok=True, mode=0o711 if traverse else 0o700)
    info = path.lstat()
    if not stat.S_ISDIR(info.st_mode) or info.st_mode & 0o022 or info.st_uid not in (0, os.getuid(), uid):
        raise WorkerBlocked("sales_codex_job_path_unsafe")
    if os.getuid() == 0:
        os.chown(path, uid, gid)
    os.chmod(path, 0o711 if traverse else 0o700)
    return path


def _write(path, raw, uid, gid):
    if path.exists() or path.is_symlink():
        raise WorkerBlocked("sales_codex_export_exists")
    descriptor = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY | os.O_NOFOLLOW, 0o600)
    with os.fdopen(descriptor, "wb") as stream:
        stream.write(raw)
        stream.flush()
        os.fsync(stream.fileno())
    if os.getuid() == 0:
        os.chown(path, uid, gid)
    os.chmod(path, 0o600)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _excel_rows(raw, mime, attachment_id, digest, customer):
    """Keep full physical cells alongside the existing bounded row classifier."""
    parsed, reports = parse_excel(raw, mime, attachment_id, digest, customer)
    physical = {}
    if mime == "application/vnd.ms-excel":
        import xlrd
        book = xlrd.open_workbook(file_contents=raw, on_demand=True)
        try:
            for index in range(book.nsheets):
                sheet = book.sheet_by_index(index)
                for number in range(sheet.nrows):
                    cells = []
                    for column in range(sheet.ncols):
                        value = sheet.cell_value(number, column)
                        if sheet.cell_type(number, column) == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).date().isoformat()
                        cells.append(_cell_text(value))
                    physical[(index, number + 1)] = (cells, False)
        finally:
            book.release_resources()
    else:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            for index, sheet in enumerate(book.worksheets):
                for number, row in enumerate(sheet.iter_rows(), 1):
                    physical[(index, number)] = ([_cell_text(cell.value) for cell in row],
                                                any(cell.data_type == "f" for cell in row))
        finally:
            book.close()
    by_id = {item["source_ref"]["id"]: item for item in parsed}
    rows, provenance = [], {}
    for report in reports:
        cells, formulas = physical[(report["sheet_index"], report["row"])]
        text = " | ".join(cells)
        nonempty = any(value.strip() for value in cells)
        ignored = report["status"] == "ignored_with_reason"
        # Cached-value extraction must not turn a formula or unknown nonempty
        # physical row into a silent blank/heading omission.
        if formulas or (ignored and report["reason"] == "blank_row" and nonempty):
            ignored = False
        item = {"id": report["id"], "source_id": attachment_id, "sha256": digest,
                "sheet": report["sheet"], "sheet_index": report["sheet_index"],
                "row": report["row"], "cells": cells, "text": text,
                "ignored": ignored, "ignore_reason": report["reason"] if ignored else "",
                "requires_review": formulas or report["reason"] in {
                    "sheet_header_unrecognized", "source_customer_mismatch"},
                "parsed": by_id.get(report["id"], {})}
        rows.append(item)
        if not ignored:
            provenance[item["id"]] = {"source_id": attachment_id, "sha256": digest,
                "sheet": item["sheet"], "sheet_index": item["sheet_index"], "row": item["row"],
                "original_text": text}
    if len(provenance) > 500:
        raise WorkerBlocked("sales_codex_source_row_limit")
    return rows, provenance


def _document(raw, mime, source_id, config):
    if config.document_python is None or config.document_python.resolve() == Path(sys.executable).resolve():
        return _bounded_extract(raw, mime, source_id)
    process = subprocess.Popen([str(config.document_python), "-m", "core.sales_document_pipeline",
        "--extract-child", mime, source_id, "null"], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL, env={"PATH": "/usr/bin:/bin", "PYTHONPATH": str(SERVICE),
            "PYTHONDONTWRITEBYTECODE": "1", "LANG": "C.UTF-8"}, start_new_session=True)
    try:
        output, _ = process.communicate(raw, timeout=35)
        if process.returncode or len(output) > 3 * 1024 * 1024:
            raise WorkerBlocked("sales_codex_document_extract_failed")
        return _extraction_from_json(runner._json(output))
    except subprocess.TimeoutExpired as error:
        os.killpg(process.pid, signal.SIGKILL)
        process.communicate()
        raise WorkerBlocked("sales_codex_document_extract_timeout") from error


def _render_page(pdf_path, page, directory, config, uid, gid):
    binary = config.pdftoppm_path
    if binary is None:
        return None
    if not binary.is_absolute() or not binary.is_file() or not os.access(binary, os.X_OK):
        return None
    prefix = directory / (pdf_path.stem + "-render-" + str(page))
    try:
        result = subprocess.run([str(binary), "-f", str(page), "-l", str(page), "-singlefile",
            "-scale-to", "1800", "-png", str(pdf_path), str(prefix)],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=25,
            env={"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8"}, check=False)
        output = prefix.with_suffix(".png")
        if result.returncode or not output.exists():
            return None
        raw = _secure_bytes(output, maximum=runner.MAX_IMAGE_BYTES)
        if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
            return None
        if os.getuid() == 0:
            os.chown(output, uid, gid)
        os.chmod(output, 0o600)
        return output, raw
    except (OSError, subprocess.TimeoutExpired, WorkerBlocked):
        return None


def _prepare_sources(store, snapshot, phase_dir, config, uid, gid):
    sources = source_manifest(snapshot)
    source_rows, provenance, attachments = [], {}, []
    forced_review, trusted_text = set(), {item["source_id"]: item.get("text", "") for item in sources}
    directory = _directory(phase_dir / "attachments", uid, gid)
    by_source = {item["source_id"]: item for item in sources}
    images, document_lines = 0, 0
    for item in snapshot["attachments"]:
        if (not re.fullmatch(r"[0-9a-f]{32}\.bin", item["blob_name"])
                or not re.fullmatch(r"[0-9a-f]{64}", item["sha256"])
                or type(item["size"]) is not int or not 0 < item["size"] <= runner.MAX_IMAGE_BYTES):
            raise WorkerBlocked("sales_codex_attachment_identity_invalid")
        blob = store.runtime / "sales-assistant" / "blobs" / item["blob_name"]
        raw = _secure_bytes(blob, maximum=runner.MAX_IMAGE_BYTES)
        if len(raw) != item["size"] or sha256(raw).hexdigest() != item["sha256"]:
            raise WorkerBlocked("sales_codex_attachment_changed")
        mime = item["mime_type"]
        extension = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp",
            "application/pdf": "pdf", "application/vnd.ms-excel": "xls",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx"}.get(mime)
        if extension is None:
            raise WorkerBlocked("sales_codex_attachment_type_unsupported")
        destination = directory / (item["id"] + "." + extension)
        _write(destination, raw, uid, gid)
        exported = {key: item[key] for key in ("id", "name", "mime_type", "sha256")}
        exported["path"] = str(destination.relative_to(phase_dir))
        attachments.append(exported)
        source = by_source[item["id"]]
        if mime in EXCEL_TYPES:
            rows, refs = _excel_rows(raw, mime, item["id"], item["sha256"], snapshot["customer_name"])
            source_rows.extend(rows); provenance.update(refs)
            source["physical_rows"] = len(rows)
            source["mandatory_rows"] = len(refs)
            trusted_text[item["id"]] = "\n".join(row["text"] for row in rows)
            if any(row["requires_review"] for row in rows):
                forced_review.add(item["id"])
        elif mime in IMAGE_TYPES:
            images += 1
            source["visual_pages"] = [{"page": 1, "attachment_path": exported["path"]}]
        elif mime == "application/pdf":
            if not raw.startswith(b"%PDF-"):
                raise WorkerBlocked("sales_codex_pdf_signature_invalid")
            try:
                extraction = _document(raw, mime, item["id"], config)
            except (ValueError, OSError):
                extraction = None
            if extraction is None:
                source.update(extraction_status="unavailable", pages=[])
                forced_review.add(item["id"])
                continue
            source["extraction_status"] = extraction.status
            source["page_count"] = extraction.page_count
            source["pages"] = []
            document_lines += len(extraction.lines)
            if document_lines > 500:
                raise WorkerBlocked("sales_codex_document_line_limit")
            by_page = {}
            for line in extraction.lines:
                by_page.setdefault(line.page_number, []).append({"line": line.line_number, "text": line.text})
            trusted_text[item["id"]] = "\n".join(line.text for line in extraction.lines)
            unready = set(extraction.unready_pages)
            if not extraction.page_count or extraction.status in {"invalid", "encrypted", "limit_exceeded", "dependency_missing"}:
                forced_review.add(item["id"])
            for page in range(1, extraction.page_count + 1):
                page_info = {"page": page, "lines": by_page.get(page, []), "visual_attached": False}
                if page in unready:
                    if images >= MAX_RENDERED_IMAGES:
                        raise WorkerBlocked("sales_codex_image_limit")
                    rendered = _render_page(destination, page, directory, config, uid, gid)
                    if rendered is None:
                        page_info["requires_review"] = True
                        forced_review.add(item["id"])
                    else:
                        image_path, image_bytes = rendered
                        visual = {"id": item["id"], "name": item["name"] + " 第" + str(page) + "页",
                            "mime_type": "image/png", "sha256": sha256(image_bytes).hexdigest(),
                            "path": str(image_path.relative_to(phase_dir))}
                        attachments.append(visual); images += 1
                        page_info.update(visual_attached=True, attachment_path=visual["path"])
                source["pages"].append(page_info)
    if images > MAX_RENDERED_IMAGES or len(provenance) > 500:
        raise WorkerBlocked("sales_codex_source_limit")
    return sources, source_rows, provenance, attachments, forced_review, trusted_text


def _safe_draft(draft):
    headers = ("customer_order_no", "delivery_date", "note", "tax_basis", "tax_rate")
    value = {key: draft.get(key, "") for key in headers}
    fields = {"id", "product_code", "quantity", "unit_price", "unit", "delivery_date",
        "source_ref", "customer_part_no", "project_no", "customer_line_no", "drawing_no",
        "product_name", "spec", "grade", "surface", "standard", "material"}
    value["lines"] = [{key: line[key] for key in fields if key in line} for line in draft.get("lines", [])]
    return value


def _packet(snapshot, *, phase, sources, rows, attachments, context, first_pass=None):
    result = {"schema_version": "ul_sales_codex_job_v1", "job_id": snapshot["job_id"],
        "session_id": snapshot["session_id"], "base_version": snapshot["base_version"],
        "profile_id": snapshot["profile_id"], "customer": snapshot["customer_name"],
        "current_data_version": snapshot["current_data_version"], "phase": phase,
        "messages": [{key: item[key] for key in ("id", "role", "text", "attachment_id")}
                     for item in snapshot["messages"]], "draft": _safe_draft(snapshot["draft"]),
        "sources": sources, "source_rows": rows, "attachments": attachments, "matching_context": context}
    if snapshot.get('previous_proposal'):
        result['previous_proposal'] = snapshot['previous_proposal']
    if first_pass is not None:
        result["first_pass"] = first_pass
    result["input_hash"] = sha256(dumps(result).encode()).hexdigest()
    if len(dumps(result).encode()) > MAX_PACKET:
        raise WorkerBlocked("sales_codex_packet_limit")
    return result


def _invoke_runner(phase_dir, config):
    # Inference uses only the stdlib, so do not expose/traverse the private
    # ulcollector dependency virtualenv from the owner's unprivileged account.
    arguments = ["/usr/bin/python3", str(SERVICE / "tools" / "sales_codex_runner.py"),
        "--job-dir", str(phase_dir), "--rules-root", str(config.rules_root),
        "--codex-binary", str(config.codex_binary), "--timeout-seconds", str(config.timeout_seconds)]
    uid, gid, account = _identity(config)
    env = {"PATH": "/usr/bin:/bin", "PYTHONDONTWRITEBYTECODE": "1", "TZ": "Asia/Shanghai", "LANG": "C.UTF-8"}
    if account is not None:
        arguments = ["/usr/sbin/runuser", "-u", account.pw_name, "--", *arguments]
        env.update(HOME=account.pw_dir, USER=account.pw_name, LOGNAME=account.pw_name)
    else:
        env.update({key: os.environ[key] for key in ("HOME", "USER", "LOGNAME") if key in os.environ})
    if os.environ.get('UL_SALES_CODEX_TRANSPORT') == 'app-server':
        env['UL_SALES_CODEX_TRANSPORT'] = 'app-server'
    process = subprocess.Popen(arguments, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                               env=env, start_new_session=True)
    try:
        import time
        deadline=time.monotonic()+config.timeout_seconds+30
        while time.monotonic()<deadline:
            if config.cancel_check and config.cancel_check():
                os.killpg(process.pid, signal.SIGTERM)
                try:process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    os.killpg(process.pid,signal.SIGKILL);process.wait(timeout=5)
                raise WorkerBlocked('sales_codex_cancelled')
            try:return process.wait(timeout=.5)
            except subprocess.TimeoutExpired:pass
        raise subprocess.TimeoutExpired(arguments,config.timeout_seconds+30)
    except subprocess.TimeoutExpired as error:
        os.killpg(process.pid, signal.SIGTERM)
        try:
            process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid, signal.SIGKILL)
            process.wait(timeout=5)
        raise WorkerBlocked("sales_codex_runner_timeout") from error


def _validated_result(phase_dir, packet, config, rules_hash, schema_hash, exit_code):
    uid, _, _ = _identity(config)
    raw_input = _secure_bytes(phase_dir / "input.json", owner=uid)
    expected_input = (dumps(packet) + "\n").encode()
    if raw_input != expected_input:
        raise WorkerBlocked("sales_codex_packet_changed")
    status = runner._json(_secure_bytes(phase_dir / "status.json", owner=uid))
    if not isinstance(status, dict):
        raise WorkerBlocked("sales_codex_runner_proof_invalid")
    expected = {key: packet[key] for key in ("job_id", "session_id", "base_version", "input_hash")}
    expected.update(input_file_sha256=sha256(expected_input).hexdigest(),
                    rules_hash=rules_hash, output_schema_sha256=schema_hash)
    if any(type(status.get(key)) is not type(value) or status.get(key) != value for key, value in expected.items()):
        raise WorkerBlocked("sales_codex_runner_binding_mismatch")
    if exit_code != 0 or status.get("status") != "succeeded":
        code = "sales_codex_runner_timed_out" if status.get("status") == "timed_out" else "sales_codex_runner_failed"
        raise WorkerBlocked(code)
    if (status.get("engine") != "codex" or status.get("external_submit") is not False
            or type(status.get("tool_calls")) is not int or status["tool_calls"] != 0):
        raise WorkerBlocked("sales_codex_runner_proof_invalid")
    result_raw = _secure_bytes(phase_dir / "result.json", owner=uid)
    if sha256(result_raw).hexdigest() != status.get("result_sha256"):
        raise WorkerBlocked("sales_codex_runner_result_changed")
    result = runner._json(result_raw)
    runner.validate_schema(result, runner._json(_secure_bytes(runner.SCHEMA_PATH)))
    return result


def _phase(snapshot, phase_dir, packet, config, uid, gid, rules_hash, schema_hash, invoke):
    _write(phase_dir / "input.json", (dumps(packet) + "\n").encode(), uid, gid)
    code = invoke(phase_dir, config)
    return _validated_result(phase_dir, packet, config, rules_hash, schema_hash, code)


def _matching_lines(first_pass, trusted_text):
    values = []
    for index, row in enumerate(first_pass["draft"]["lines"], 1):
        code = row["product_code"]
        raw = "\n".join(trusted_text.get(identity, "") for identity in row["source_ids"])
        explicit = bool(code and re.search(r"(?<![A-Za-z0-9_.-])" + re.escape(code) + r"(?![A-Za-z0-9_.-])", raw))
        reference = {"id": "codex-line-" + str(index), "original_name": row["original_name"],
            "original_spec": row["original_spec"], "original_standard": row["standard"],
            "original_grade": row["grade"], "original_surface": row["surface"],
            "original_material": row["original_material"],
            "customer_part_no": row["customer_part_no"], "source_ids": row["source_ids"]}
        values.append({"source_ref": reference, "product_code": code if explicit else ""})
    return values


def _review_result(store, snapshot, result, context, rows, provenance, sources, forced, trusted_text):
    source_ids = [item["source_id"] for item in sources]
    # validate_output may normalize numeric fields; give it only an isolated copy.
    validate_output(json.loads(dumps(result)), source_ids, provenance.keys())
    expected_rows = set(provenance)
    occurrences = {identity: 0 for identity in expected_rows}
    evidence_rows = [(line, True) for line in result["draft"]["lines"]]
    evidence_rows.extend((line, False) for line in result["new_product_suggestions"])
    for line, order_line in evidence_rows:
        quoted = re.sub(r"\s+", "", line["source_quote"])
        originals = [re.sub(r"\s+", "", trusted_text.get(identity, "")) for identity in line["source_ids"]]
        visual = any(item["source_id"] in line["source_ids"] and
            (item.get("visual_pages") or any(page.get("visual_attached") for page in item.get("pages", [])))
            for item in sources)
        if not quoted or (not visual and not any(quoted in raw for raw in originals if raw)):
            raise WorkerBlocked("sales_codex_source_quote_mismatch")
        refs = set(line["source_ids"]) & expected_rows
        if order_line and len(refs) > 1:
            raise WorkerBlocked("sales_codex_rows_merged")
        if order_line:
            for identity in refs:
                occurrences[identity] += 1
    if any(count > 1 for count in occurrences.values()):
        raise WorkerBlocked("sales_codex_source_row_repeated")
    for entry in result["source_coverage"]:
        if entry["source_id"] in forced:
            entry.update(status="needs_review", detail="服务端发现公式、未知行或文档提取不完整，须核对原件")
    if forced:
        result["status"] = "needs_clarification"
        result["questions"].append("原资料还有未验证的公式、数据行或PDF页，请补充或核对原件。")
    candidates = context["matching"]["lines"]
    ambiguous = []
    existing_sources = set()
    with store.connect() as db:
        for number, (row, evidence) in enumerate(zip(result["draft"]["lines"], candidates), 1):
            if evidence.get('available_candidate_count', 0):
                existing_sources.update(row['source_ids'])
            code = row["product_code"]
            if not code:
                continue
            offered = [candidate for candidate in evidence["candidates"] if candidate["product_code"] == code]
            if not offered:
                raise WorkerBlocked("sales_codex_product_not_offered")
            if offered and any(not candidate.get("available") for candidate in offered):
                result["status"] = "needs_clarification"
                result["questions"].append("商品 " + code + " 的匹配证据仍有待核对项目。")
            elif evidence.get('needs_question'):
                ambiguous.append(str(number))
    if ambiguous:
        result['status'] = 'needs_clarification'
        if not result['questions']:
            result['questions'].append('第' + '、'.join(ambiguous) + '行存在多个现有编码或缺少匹配属性，请核对候选差异。')
    # A model must not send an existing candidate into the new-product flow.
    removed = [item for item in result['new_product_suggestions']
               if set(item['source_ids']) & existing_sources]
    if removed:
        result['new_product_suggestions'] = [item for item in result['new_product_suggestions'] if item not in removed]
        result['status'] = 'needs_clarification'
        result['questions'].append('这些明细已有可用商品候选，请从现有编码中核对，不需要新建商品。')
    if not context["matching"].get("catalog_complete"):
        result["status"] = "needs_clarification"
        result["questions"].append("当前商品目录覆盖未通过，不能直接采用商品匹配结果。")
    # Keep every line/evidence while presenting a bounded, deduplicated list.
    # A large batch with unavailable candidates must remain clarifiable rather
    # than failing the downstream JSON contract due to repeated questions.
    questions = list(dict.fromkeys(result["questions"]))
    if len(questions) > 30:
        questions = questions[:29] + ["另有 " + str(len(questions) - 29) + " 项待核对，请查看逐行匹配证据。"]
    result["questions"] = questions
    return result


def run_one(store, config, *, history_reference=None, invoke_runner=_invoke_runner):
    store = _ServiceIdentityStore(store)
    pipeline = SalesCodexPipeline(store)
    token = uuid.uuid4().hex
    snapshot = pipeline.claim(token)
    if snapshot is None:
        return {"claimed": False, "external_submit": False}
    job_id = snapshot["job_id"]
    def cancelled():
        with store.connect() as db:
            row=db.execute('SELECT status,worker_token FROM sales_codex_jobs WHERE id=?',(job_id,)).fetchone()
            return not row or row['status']!='running' or row['worker_token']!=token
    config=replace(config,cancel_check=cancelled)
    try:
        uid, gid, _ = _identity(config)
        _directory(config.jobs_root, os.getuid(), os.getgid(), traverse=True)
        job_root = _directory(config.jobs_root / job_id, os.getuid(), os.getgid(), traverse=True)
        first_dir = job_root / "phase-1"
        if first_dir.exists():
            raise WorkerBlocked("sales_codex_job_export_exists")
        _directory(first_dir, uid, gid)
        import shutil
        base_text, base_hash = runner._rules(config.rules_root)
        version=snapshot.get('rules_version',{})
        if version.get('base_hash') and version['base_hash']!=base_hash:
            raise WorkerBlocked('sales_rules_base_changed')
        pinned_root=job_root/'rules'
        shutil.copytree(config.rules_root,pinned_root,ignore=shutil.ignore_patterns('.*'))
        if version.get('content'):
            if sha256(version['content'].encode()).hexdigest()!=version.get('sha256'):
                raise WorkerBlocked('sales_rules_version_changed')
            (pinned_root/'references'/'owner-supplement.md').write_text(version['content'])
        for path in [pinned_root,*pinned_root.rglob('*')]:
            path.chmod(0o755 if path.is_dir() else 0o644)
        config=replace(config,rules_root=pinned_root)
        _, rules_hash = runner._rules(config.rules_root)
        schema_hash = sha256(_secure_bytes(runner.SCHEMA_PATH)).hexdigest()
        rules = load_skill_rules(config.rules_root / "references" / "rules")
        sources, rows, provenance, attachments, forced, trusted_text = _prepare_sources(
            store, snapshot, first_dir, config, uid, gid)
        first_packet = _packet(snapshot, phase="analyze", sources=sources, rows=rows,
                               attachments=attachments, context={})
        first = _phase(snapshot, first_dir, first_packet, config, uid, gid,
                       rules_hash, schema_hash, invoke_runner)
        source_ids = [item["source_id"] for item in sources]
        validate_output(json.loads(dumps(first)), source_ids, provenance.keys())
        pipeline.progress(job_id, token, "match")
        lines = _matching_lines(first, {**trusted_text, **{key: value["original_text"] for key, value in provenance.items()}})
        history = history_reference or SalesHistoryReference()
        with history.connect() as reference, store.connect() as db:
            job = db.execute("SELECT * FROM sales_codex_jobs WHERE id=?", (job_id,)).fetchone()
            session = db.execute("SELECT * FROM sales_assistant_sessions WHERE id=?", (snapshot["session_id"],)).fetchone()
            if not pipeline._current(db, job, session):
                raise WorkerBlocked("sales_codex_stale")
            matching = build_matching_context(db, snapshot["profile_id"], snapshot["customer_name"],
                                              lines, history_db=reference, rules=rules)
        context = {"matching": matching, "source_lines": lines}
        second_dir = _directory(job_root / "phase-2", uid, gid)
        _directory(second_dir / "attachments", uid, gid)
        for item in attachments:
            raw = _secure_bytes(first_dir / item["path"], maximum=runner.MAX_IMAGE_BYTES, owner=uid)
            if sha256(raw).hexdigest() != item["sha256"]:
                raise WorkerBlocked("sales_codex_export_changed")
            _write(second_dir / item["path"], raw, uid, gid)
        pipeline.progress(job_id, token, "review")
        second_packet = _packet(snapshot, phase="match_and_review", sources=sources, rows=rows,
                                attachments=attachments, context=context, first_pass=first)
        final = _phase(snapshot, second_dir, second_packet, config, uid, gid,
                       rules_hash, schema_hash, invoke_runner)
        all_text = {**trusted_text, **{key: value["original_text"] for key, value in provenance.items()}}
        validate_output(json.loads(dumps(final)), source_ids, provenance.keys())
        # Rebuild each final line independently; the model cannot swap source
        # attributes/SKU after the first-phase evidence was supplied.
        final_lines = _matching_lines(final, all_text)
        with history.connect() as reference, store.connect() as db:
            matching_final = build_matching_context(db, snapshot["profile_id"], snapshot["customer_name"],
                final_lines, history_db=reference, rules=rules)
            dependencies = product_dependencies(db, snapshot['profile_id'],
                [line['product_code'] for line in final['draft']['lines']])
        final_context = {"source_lines": final_lines, "matching": matching_final}
        final = _review_result(store, snapshot, final, final_context, rows, provenance, sources, forced, all_text)
        runner._atomic_json(job_root / "matching-proof.json", {"job_id": job_id,
            "base_version": snapshot["base_version"], "data_version": snapshot["current_data_version"],
            "rules_hash": rules_hash, "verified_phase_result_hash": sha256(dumps(final).encode()).hexdigest(),
            "matching": matching_final, "external_submit": False})
        result = pipeline.finish(job_id, token, output=final, source_ids=source_ids,
                                 physical_rows=list(provenance), source_provenance=provenance,
                                 catalog_dependencies=dependencies)
        return {"claimed": True, **result, "external_submit": False}
    except Exception as error:
        if isinstance(error, (WorkerBlocked, CodexJobError)):
            code = error.code
        elif isinstance(error, runner.RunnerBlocked):
            code = "sales_codex_" + str(error)
        else:
            code = "sales_codex_worker_failed"
        try:
            result = pipeline.finish(job_id, token, error_code=code)
        except CodexJobError:
            result = {"id": job_id, "status": "failed", "error_code": "sales_codex_claim_changed"}
        return {"claimed": True, **result, "error_code": code, "external_submit": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runtime", required=True)
    parser.add_argument("--jobs-root", required=True, type=Path)
    parser.add_argument("--rules-root", required=True, type=Path)
    parser.add_argument("--codex-binary", type=Path, default=Path("/home/ulcodex/.local/bin/codex"))
    parser.add_argument("--codex-user", default="ulcodex")
    parser.add_argument("--timeout-seconds", type=int, default=180)
    parser.add_argument("--document-python", type=Path)
    parser.add_argument("--pdftoppm-path", type=Path)
    parser.add_argument("--history-reference", default=os.environ.get("UL_SALES_HISTORY_REFERENCE_PATH", FORMAL_HISTORY))
    parser.add_argument("--history-sha256", default=os.environ.get("UL_SALES_HISTORY_REFERENCE_SHA256", FORMAL_HISTORY_SHA))
    args = parser.parse_args()
    try:
        reference = SalesHistoryReference(args.history_reference, args.history_sha256)
        config = WorkerConfig(args.jobs_root, args.rules_root, args.codex_binary, args.codex_user,
                              args.timeout_seconds, args.document_python, args.pdftoppm_path)
        database_path = Path(args.runtime) / "workbench.sqlite3"
        info = database_path.stat()
        with _database_identity(info.st_uid, info.st_gid):
            store = Store(args.runtime)
        result = run_one(store, config, history_reference=reference)
        if not result.get("claimed"):
            from tools.sales_rule_eval_worker import run_one as evaluate_rules
            result=evaluate_rules(_ServiceIdentityStore(store),config)
        print(dumps(result))
        return 0
    except (ValueError, OSError):
        print(dumps({"claimed": False, "status": "failed", "code": "sales_codex_worker_config_invalid", "external_submit": False}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
