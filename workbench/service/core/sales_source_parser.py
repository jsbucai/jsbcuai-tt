"""Bounded, deterministic source-row extraction for sales drafts.

This module never matches a product or submits a document. Every physical
source row receives a report entry, including blank, heading, and footer rows.
Only an explicit product code can become an automatic product candidate.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from io import StringIO, BytesIO
import re
import unicodedata


MAX_SHEETS = 8
MAX_PHYSICAL_ROWS = 600
MAX_DATA_ROWS = 500
MAX_COLUMNS = 60

_ALIASES = {
    'product_code': ('商品编号', '商品编码', '编码', '产品编码', 'ERP商品编号', '系统编码'),
    'product_name': ('物料名称', '物资名称', '商品名称', '产品名称', '品名', '中文描述'),
    'customer_part_no': ('客户料号', '客户品号', '客户商品编号', '物料号', '物料编号', '料号', '图号', '图纸编号', '物料'),
    'customer_order_no': ('客户订单号', '客户单号', '客户PO', '采购订单号', '订单号'),
    'project_no': ('项目号', '项目编号'),
    'customer_line_no': ('客户行号', '客户行次', '行号', '项次', '序号'),
    'standard_no': ('标准', '标准号', '执行标准'),
    'quantity': ('数量', '订货数量', '订单数量', '订购数量', '采购数量', '交货数量', '实际数', '数量1'),
    'unit_price': ('单价', '销售单价', '报价'),
    'taxed_price': ('含税单价',),
    'untaxed_price': ('未税单价', '不含税单价'),
    'unit': ('单位', '计量单位'),
    'warehouse': ('仓库', '仓库编号'),
    'delivery_date': ('交期', '交货日期', '到货日期', '订单交期', '计划交货日期'),
    'customer_name': ('客户名称', '往来单位名称'),
    'spec': ('规格', '物资规格', '规格型号', '规格型号_原始', '标准号/规格'),
    'grade': ('等级', '性能等级'),
    'surface': ('表面处理', '材质/表面处理'),
}
_HEADER_LOOKUP = {name: field for field, names in _ALIASES.items() for name in names}
_CODE = re.compile(r'[A-Za-z0-9_.-]+\Z')
_PART_LABEL = re.compile(r'^(?:客户料号|客户品号|图号|料号)\s*[:：=]\s*(.+)$')
_NUMERIC_FIELD = re.compile(r'[-+]?\d+(?:\.\d+)?\Z')
_FOOTER = frozenset({'合计', '小计', '总计'})


def _text(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _header_name(value: object) -> str:
    name = re.sub(r'\s+', '', unicodedata.normalize('NFKC', _text(value)))
    # Strip only an explicit price-unit annotation, never arbitrary qualifiers
    # such as “未税单价(暂定)” or a second price column.
    return re.sub(r'(?<=单价)\((?:元|人民币|RMB)/(?:件|支|个)\)$', '', name)


def _preview(values: object) -> str:
    if isinstance(values, (list, tuple)):
        result = ' | '.join(_text(value) for value in values if _text(value))
    else:
        result = str(values).strip()
    return result[:240]


def _report(source_kind: str, source_id: str, digest: str, row: int,
            preview: str, *, sheet: str = '', sheet_index: int = 0,
            reason: str = '', ignored: bool = False) -> dict:
    identity = f'{source_kind}:{source_id}:{sheet_index}:{row}'
    return {'id': identity, 'kind': source_kind, 'source_id': source_id,
            'sha256': digest, 'sheet': sheet, 'sheet_index': sheet_index,
            'row': row, 'preview': preview, 'status': 'ignored_with_reason' if ignored else 'needs_review',
            'reason': reason, 'blocking': not ignored, 'line_id': ''}


def _ref(report: dict, *, name: str = '', spec: str = '', grade: str = '', surface: str = '',
         customer_part_no: str = '', customer_order_no: str = '', project_no: str = '',
         customer_line_no: str = '', standard_no: str = '', unit_blank: bool = False,
         explicit_code: bool = False) -> dict:
    return {'id': report['id'], 'kind': report['kind'], 'source_id': report['source_id'],
            'sha256': report['sha256'], 'sheet': report['sheet'], 'row': report['row'],
            'preview': report['preview'], 'original_name': name, 'original_spec': spec,
            'original_grade': grade, 'original_surface': surface,
            'customer_part_no': customer_part_no, 'customer_order_no': customer_order_no,
            'project_no': project_no, 'customer_line_no': customer_line_no,
            'original_standard': standard_no,
            'unit_blank': unit_blank,
            'match_basis': 'explicit_code' if explicit_code else 'manual_code_required'}


def parse_text(text: str, message_id: str, digest: str) -> tuple[list[dict], list[dict]]:
    """Parse two/three-column CSV or whitespace rows; keep other lines for review."""
    lines, reports = [], []
    for row_number, raw in enumerate(text.splitlines(), start=1):
        if row_number > MAX_DATA_ROWS:
            raise ValueError('文字明细最多 500 行')
        preview = _preview(raw)
        if not preview:
            reports.append(_report('text', message_id, digest, row_number, '',
                                   reason='blank_row', ignored=True))
            continue
        try:
            values = next(csv.reader(StringIO(raw), strict=True))
        except csv.Error:
            values = [raw]
        if len(values) == 1 and ',' not in raw:
            values = raw.split()
        values = [part.strip() for part in values]
        explicit = len(values) in (2, 3) and bool(_CODE.fullmatch(values[0]))
        part = _PART_LABEL.fullmatch(values[0]) if values else None
        structured = explicit or (len(values) in (2, 3) and
            (part is not None or bool(_NUMERIC_FIELD.fullmatch(values[1])))) or (
            len(values) == 4 and (part is not None or not _CODE.fullmatch(values[0]))
            and bool(_NUMERIC_FIELD.fullmatch(values[2])))
        source_spec = values[1] if structured and len(values) == 4 else ''
        quantity_index = 2 if len(values) == 4 else 1
        report = _report('text', message_id, digest, row_number, preview,
                         reason='' if explicit else 'source_line_needs_match')
        reports.append(report)
        item = {'product_code': values[0] if explicit else '',
                'quantity': values[quantity_index] if structured else '',
                'unit_price': values[quantity_index + 1] if structured and len(values) > quantity_index + 1 else '',
                'source_ref': _ref(report, explicit_code=explicit,
                    customer_part_no=part.group(1).strip() if part is not None else '',
                    name=values[0] if structured and not explicit and part is None else '',
                    spec=source_spec)}
        lines.append(item)
    return lines, reports


def _columns(values: tuple) -> dict[str, int] | None:
    found = {}
    for index, value in enumerate(values):
        field = _HEADER_LOOKUP.get(_header_name(value))
        if field:
            if field in found:
                raise ValueError('Excel 表头字段重复，请人工核对')
            found[field] = index
    if 'quantity' not in found or not ({'product_code', 'product_name', 'customer_part_no'} & found.keys()):
        return None
    price_columns = {'unit_price', 'taxed_price', 'untaxed_price'} & found.keys()
    if len(price_columns) > 1:
        raise ValueError('Excel 请仅保留一列单价，并注明含税或未税')
    return found


def _sheet_rows(rows: list[tuple[int, tuple]], sheet: str, sheet_index: int,
                source_id: str, digest: str, customer: str,
                date_cell) -> tuple[list[dict], list[dict]]:
    candidates = []
    for index, (_number, values) in enumerate(rows[:30]):
        columns = _columns(values)
        if columns is not None:
            candidates.append((index, columns))
    if len(candidates) > 1:
        raise ValueError(f'Excel 工作表 {sheet} 存在多个候选表头，请人工核对')
    header_index, columns = candidates[0] if candidates else (None, None)
    parsed, reports = [], []
    for index, (row_number, values) in enumerate(rows):
        preview = _preview(values)
        if not preview:
            reports.append(_report('spreadsheet', source_id, digest, row_number, '',
                                   sheet=sheet, sheet_index=sheet_index,
                                   reason='blank_row', ignored=True))
            continue
        if header_index is None:
            reports.append(_report('spreadsheet', source_id, digest, row_number, preview,
                                   sheet=sheet, sheet_index=sheet_index,
                                   reason='sheet_header_unrecognized'))
            continue
        if index <= header_index:
            reports.append(_report('spreadsheet', source_id, digest, row_number, preview,
                                   sheet=sheet, sheet_index=sheet_index,
                                   reason='header' if index == header_index else 'before_header',
                                   ignored=True))
            continue
        cells = {field: date_cell(row_number, column, values[column])
                 if field == 'delivery_date' and column < len(values) else
                 values[column] if column < len(values) else ''
                 for field, column in columns.items()}
        code = _text(cells.get('product_code'))
        first = next((_text(value) for value in values if _text(value)), '')
        if first in _FOOTER and not code:
            reports.append(_report('spreadsheet', source_id, digest, row_number, preview,
                                   sheet=sheet, sheet_index=sheet_index,
                                   reason='footer', ignored=True))
            continue
        name = _text(cells.get('product_name'))
        quantity = _text(cells.get('quantity'))
        source_customer = _text(cells.get('customer_name'))
        if source_customer and source_customer != customer:
            reports.append(_report('spreadsheet', source_id, digest, row_number, preview,
                                   sheet=sheet, sheet_index=sheet_index,
                                   reason='source_customer_mismatch'))
            continue
        price_field = next((field for field in ('taxed_price', 'untaxed_price', 'unit_price')
                            if field in columns), '')
        explicit = bool(code and _CODE.fullmatch(code))
        report = _report('spreadsheet', source_id, digest, row_number, preview,
                         sheet=sheet, sheet_index=sheet_index,
                         reason='' if explicit else 'source_line_needs_match')
        reports.append(report)
        item = {'product_code': code if explicit else '', 'quantity': quantity,
                'unit_price': _text(cells.get(price_field)) if price_field else '',
                'unit': _text(cells.get('unit')),
                'warehouse': _text(cells.get('warehouse')),
                'delivery_date': _text(cells.get('delivery_date')),
                'source_ref': _ref(report, name=name, spec=_text(cells.get('spec')),
                                   grade=_text(cells.get('grade')), surface=_text(cells.get('surface')),
                                   customer_part_no=_text(cells.get('customer_part_no')),
                                   customer_order_no=_text(cells.get('customer_order_no')),
                                   project_no=_text(cells.get('project_no')),
                                   customer_line_no=_text(cells.get('customer_line_no')),
                                   standard_no=_text(cells.get('standard_no')),
                                   unit_blank=not _text(cells.get('unit')),
                                   explicit_code=explicit)}
        if price_field == 'taxed_price':
            item['tax_basis'] = 'tax_inclusive'
        elif price_field == 'untaxed_price':
            item['tax_basis'] = 'tax_exclusive'
        parsed.append(item)
    return parsed, reports


def parse_excel(raw: bytes, mime: str, attachment_id: str, digest: str,
                customer: str) -> tuple[list[dict], list[dict]]:
    """Extract all sheets; never discard a nonempty row without a report."""
    parsed, reports, physical = [], [], 0
    if mime == 'application/vnd.ms-excel':
        import xlrd
        book = xlrd.open_workbook(file_contents=raw, on_demand=True)
        try:
            if book.nsheets > MAX_SHEETS:
                raise ValueError('Excel 工作表超过 8 张，请人工核对')
            for sheet_index in range(book.nsheets):
                sheet = book.sheet_by_index(sheet_index)
                if sheet.ncols > MAX_COLUMNS:
                    raise ValueError('Excel 列数超过 60 列，请人工核对')
                physical += sheet.nrows
                if physical > MAX_PHYSICAL_ROWS:
                    raise ValueError('Excel 来源行超过 600 行，请拆分文件')
                rows = [(number + 1, tuple(sheet.row_values(number))) for number in range(sheet.nrows)]
                def date_cell(number, column, value):
                    if sheet.cell_type(number - 1, column) == xlrd.XL_CELL_DATE:
                        return xlrd.xldate_as_datetime(value, book.datemode).date().isoformat()
                    return value
                items, seen = _sheet_rows(rows, sheet.name, sheet_index, attachment_id,
                                          digest, customer, date_cell)
                parsed.extend(items); reports.extend(seen)
        finally:
            book.release_resources()
    else:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise ValueError('Excel 工作表超过 8 张，请人工核对')
            for sheet_index, sheet in enumerate(book.worksheets):
                if sheet.max_column > MAX_COLUMNS:
                    raise ValueError('Excel 列数超过 60 列，请人工核对')
                if sheet.max_row > MAX_PHYSICAL_ROWS:
                    raise ValueError('Excel 来源行超过 600 行，请拆分文件')
                rows = [(number, tuple(values)) for number, values in
                        enumerate(sheet.iter_rows(values_only=True), start=1)]
                physical += len(rows)
                if physical > MAX_PHYSICAL_ROWS:
                    raise ValueError('Excel 来源行超过 600 行，请拆分文件')
                items, seen = _sheet_rows(rows, sheet.title, sheet_index, attachment_id,
                                          digest, customer, lambda _row, _col, value: value)
                parsed.extend(items); reports.extend(seen)
        finally:
            book.close()
    if len(parsed) > MAX_DATA_ROWS:
        raise ValueError('Excel 明细最多 500 行')
    return parsed, reports
