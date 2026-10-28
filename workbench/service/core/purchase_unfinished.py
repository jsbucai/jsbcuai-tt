"""Read a validated, immutable purchase export without importing ERP facts."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
from threading import RLock
import uuid
import warnings

from flask import Blueprint, g, jsonify, request
from openpyxl import load_workbook

from .daily_work import DailyWorkConflict
from .store import timestamp


SHEET = "进货订单统计列表"
REQUIRED = frozenset(("行号", "单位编号", "单位名称", "商品编号", "商品名称",
                      "订单编号", "仓库", "订货日期", "到货日期", "订货数量",
                      "完成数量", "终止完成数量", "未完成数量"))
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
SNAPSHOT_ID = re.compile(r"^[0-9]{8}T[0-9]{6}Z-[0-9a-f]{12}-[0-9a-f]{8}$")
SHA256 = re.compile(r"^[0-9a-f]{64}$")


class PurchaseSnapshotError(RuntimeError):
    def __init__(self, code: str, unavailable: bool = False):
        super().__init__(code)
        self.code, self.unavailable = code, unavailable


def _date(value, code: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise PurchaseSnapshotError(code)


def _quantity(value, code: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise PurchaseSnapshotError(code)
    try:
        number = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError):
        raise PurchaseSnapshotError(code) from None
    if not number.is_finite():
        raise PurchaseSnapshotError(code)
    return number


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def _file_stamp(path: Path, limit: int):
    try:
        value = path.lstat()
    except OSError:
        raise PurchaseSnapshotError("snapshot_file_missing_or_symlink") from None
    if not stat.S_ISREG(value.st_mode):
        raise PurchaseSnapshotError("snapshot_file_missing_or_symlink")
    if value.st_size <= 0 or value.st_size > limit:
        raise PurchaseSnapshotError("snapshot_file_size_invalid")
    return (value.st_dev, value.st_ino, value.st_size, value.st_mtime_ns, value.st_ctime_ns)


def _private_file(path: Path, limit: int) -> bytes:
    before = _file_stamp(path, limit)
    try:
        with path.open('rb') as stream:
            value = os.fstat(stream.fileno())
            opened = (value.st_dev, value.st_ino, value.st_size, value.st_mtime_ns, value.st_ctime_ns)
            if opened != before:
                raise PurchaseSnapshotError("snapshot_file_changed")
            raw = stream.read(limit + 1)
    except OSError:
        raise PurchaseSnapshotError("snapshot_file_changed") from None
    if _file_stamp(path, limit) != before or len(raw) != before[2]:
        raise PurchaseSnapshotError("snapshot_file_changed")
    return raw


def _json_file(path: Path) -> dict:
    try:
        value = json.loads(_private_file(path, 128 * 1024))
    except (UnicodeDecodeError, ValueError):
        raise PurchaseSnapshotError("snapshot_json_invalid") from None
    if not isinstance(value, dict):
        raise PurchaseSnapshotError("snapshot_json_invalid")
    return value


def _line_key(values: dict, account_scope: str) -> str:
    # The report has no ERP line ID. This is only a local reply anchor; any
    # duplicated supplier/order/product/warehouse combination stays ambiguous.
    fields = [account_scope or "legacy_unscoped_review",
              *(values[key] for key in ("supplier_code", "order_no", "product_code",
                                         "warehouse", "order_date"))]
    return sha256(json.dumps(fields, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()[:32]


def _parse_xlsx(raw: bytes, manifest: dict, account_scope: str) -> tuple[list[dict], dict]:
    start = _date(manifest.get("query_start"), "query_start_invalid")
    end = _date(manifest.get("query_end"), "query_end_invalid")
    if start > end or manifest.get("filter") != "未完成数量 >= 0.1":
        raise PurchaseSnapshotError("query_scope_invalid")
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style")
            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise PurchaseSnapshotError("workbook_invalid") from None
    try:
        if book.sheetnames != [SHEET]:
            raise PurchaseSnapshotError("workbook_sheet_invalid")
        sheet = book.active
        # The vendor publishes dimension A1 even when the report has 90+ rows.
        sheet.reset_dimensions()
        header = None
        rows = []
        suppliers, orders = set(), set()
        supplier_names = {}
        duplicates = Counter()
        unfinished_sum = Decimal(0)
        minimum = None
        total_seen = False
        for row_no, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_no > MAX_ROWS + 20:
                raise PurchaseSnapshotError("workbook_too_many_rows")
            if row_no == 1 and (not cells or cells[0] != SHEET):
                raise PurchaseSnapshotError("workbook_title_invalid")
            if row_no in (2, 3):
                expected_label = "开始时间" if row_no == 2 else "结束时间"
                expected_date = start if row_no == 2 else end
                if len(cells) < 2 or cells[0] != expected_label or _date(cells[1], "query_date_invalid") != expected_date:
                    raise PurchaseSnapshotError("query_date_mismatch")
            if row_no < 4:
                continue
            if header is None:
                if "订单编号" not in cells or "未完成数量" not in cells:
                    if row_no >= 15:
                        raise PurchaseSnapshotError("workbook_header_missing")
                    continue
                header = tuple(cells)
                if (any(not isinstance(value, str) or not value.strip() for value in header) or
                        len(header) != len(set(header)) or not REQUIRED.issubset(header) or
                        list(header) != manifest.get("columns") or row_no != manifest.get("header_row")):
                    raise PurchaseSnapshotError("workbook_header_invalid")
                index = {name: header.index(name) for name in header}
                continue
            if all(value in (None, "") for value in cells):
                continue
            values = {name: cells[position] if position < len(cells) else None
                      for name, position in index.items()}
            serial = values["行号"]
            if serial == "合计":
                if total_seen or _quantity(values["未完成数量"], "total_quantity_invalid") != unfinished_sum:
                    raise PurchaseSnapshotError("workbook_total_invalid")
                total_seen = True
                continue
            if total_seen or isinstance(serial, bool) or not isinstance(serial, (int, float)) or serial != len(rows) + 1:
                raise PurchaseSnapshotError("workbook_serial_invalid")
            for name in ("单位编号", "单位名称", "商品编号", "商品名称", "订单编号", "仓库"):
                if not _text(values[name]):
                    raise PurchaseSnapshotError("workbook_identity_missing")
            order_date = _date(values["订货日期"], "order_date_invalid")
            if not start <= order_date <= end:
                raise PurchaseSnapshotError("order_date_outside_scope")
            arrival_date = _date(values["到货日期"], "arrival_date_invalid").isoformat() if values["到货日期"] not in (None, "") else ""
            unfinished = _quantity(values["未完成数量"], "unfinished_quantity_invalid")
            if unfinished < Decimal("0.1"):
                raise PurchaseSnapshotError("unfinished_below_threshold")
            quantities = {}
            for source, target in (("订货数量", "ordered_quantity"), ("完成数量", "completed_quantity"),
                                   ("终止完成数量", "terminated_quantity")):
                quantity = _quantity(values[source], source)
                if quantity < 0:
                    raise PurchaseSnapshotError("negative_quantity")
                quantities[target] = str(quantity)
            row = {"line_number": len(rows) + 1,
                   "supplier_code": _text(values["单位编号"]), "supplier_name": _text(values["单位名称"]),
                   "product_code": _text(values["商品编号"]), "product_name": _text(values["商品名称"]),
                   "order_no": _text(values["订单编号"]), "warehouse": _text(values["仓库"]),
                   "order_date": order_date.isoformat(), "arrival_date": arrival_date,
                   "unfinished_quantity": str(unfinished), "unit": "", **quantities}
            row["line_key"] = _line_key(row, account_scope)
            duplicate_key = tuple(row[name] for name in ("supplier_code", "order_no", "product_code", "warehouse"))
            duplicates[duplicate_key] += 1
            suppliers.add(row["supplier_code"])
            old_name = supplier_names.setdefault(row["supplier_code"], row["supplier_name"])
            if old_name != row["supplier_name"]:
                raise PurchaseSnapshotError("supplier_identity_conflict")
            orders.add(row["order_no"])
            unfinished_sum += unfinished
            minimum = unfinished if minimum is None else min(minimum, unfinished)
            rows.append(row)
        if not header or not total_seen or not rows:
            raise PurchaseSnapshotError("workbook_incomplete")
        if len(rows) != manifest.get("row_count") or len(rows) != manifest.get("web_row_count"):
            raise PurchaseSnapshotError("row_count_mismatch")
        checks = {"order_count": len(orders), "supplier_count": len(suppliers),
                  "duplicate_candidate_key_groups": sum(count > 1 for count in duplicates.values()),
                  "minimum_unfinished": str(minimum), "unfinished_sum": str(unfinished_sum)}
        if any(str(manifest.get(key)) != str(value) for key, value in checks.items()):
            raise PurchaseSnapshotError("manifest_totals_mismatch")
        for row in rows:
            key = tuple(row[name] for name in ("supplier_code", "order_no", "product_code", "warehouse"))
            row["identity_ambiguous"] = duplicates[key] > 1
        return rows, checks
    finally:
        book.close()


def ensure_schema(store):
    with store.connect(write=True) as db:
        db.execute("""CREATE TABLE IF NOT EXISTS purchase_supplier_promises(
            line_key TEXT PRIMARY KEY, supplier_code TEXT NOT NULL, order_no TEXT NOT NULL,
            product_code TEXT NOT NULL, warehouse TEXT NOT NULL, order_date TEXT NOT NULL,
            promised_date TEXT NOT NULL, note TEXT NOT NULL, version INTEGER NOT NULL,
            updated_by TEXT NOT NULL, source_snapshot_id TEXT NOT NULL,
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
            account_scope TEXT NOT NULL DEFAULT '')""")
        columns = {row["name"] for row in db.execute(
            "PRAGMA table_info(purchase_supplier_promises)")}
        if "account_scope" not in columns:
            # Preserve old notes for audit, but never guess their tenant.
            db.execute("ALTER TABLE purchase_supplier_promises ADD COLUMN account_scope TEXT NOT NULL DEFAULT ''")
        db.execute("""CREATE INDEX IF NOT EXISTS purchase_promises_account_scope
            ON purchase_supplier_promises(account_scope,supplier_code)""")
        db.execute("""CREATE TABLE IF NOT EXISTS purchase_supplier_promise_events(
            id TEXT PRIMARY KEY, line_key TEXT NOT NULL, idempotency_key TEXT NOT NULL,
            actor_id TEXT NOT NULL, promised_date TEXT NOT NULL, note TEXT NOT NULL,
            source_snapshot_id TEXT NOT NULL, created_at TEXT NOT NULL,
            UNIQUE(line_key,idempotency_key))""")


class PurchaseUnfinished:
    def __init__(self, store, snapshot_dir=None, preview_dir=None):
        self.store = store
        self.root = Path(snapshot_dir) if snapshot_dir else None
        self.preview_root = Path(preview_dir) if preview_dir else None
        self._parse_lock = RLock()
        self._parsed_snapshot = None
        ensure_schema(store)

    def _validated_rows(self, path, manifest, account_scope):
        """Reuse only a complete parse whose file identity and SHA-256 passed.

        Stat includes inode and ctime, so replacing a file or restoring its
        size/mtime cannot reuse a previous parse. Failed publications leave the
        valid cache intact but still raise; they never masquerade as success.
        """
        with self._parse_lock:
            stamp = _file_stamp(path, MAX_BYTES)
            key = (str(path.absolute()), stamp, account_scope,
                   sha256(json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()).hexdigest())
            if self._parsed_snapshot and self._parsed_snapshot[0] == key:
                if _file_stamp(path, MAX_BYTES) != stamp:
                    raise PurchaseSnapshotError('snapshot_file_changed')
                return self._parsed_snapshot[1]
            raw = _private_file(path, MAX_BYTES)
            if len(raw) != manifest['bytes'] or sha256(raw).hexdigest() != manifest['sha256']:
                raise PurchaseSnapshotError('snapshot_hash_or_size_mismatch')
            rows, _ = _parse_xlsx(raw, manifest, account_scope)
            if _file_stamp(path, MAX_BYTES) != stamp:
                raise PurchaseSnapshotError('snapshot_file_changed')
            self._parsed_snapshot = (key, rows)
            return rows

    @staticmethod
    def _authorize(user):
        if not isinstance(user, dict) or not user.get("active", True):
            raise DailyWorkConflict("请先登录采购岗位", "role_forbidden")
        if not user.get("is_admin") and "purchase" not in user.get("roles", []):
            raise DailyWorkConflict("账号没有采购岗位权限", "role_forbidden")

    def _load(self) -> tuple[dict, list[dict]]:
        root = self.root
        stage = "formal"
        # A validated collector trial can be viewed on the phone while its
        # business publication gate is still closed.  The formal pointer, if
        # present, always takes precedence; an invalid formal snapshot must
        # not silently fall back to trial data.
        if (root is None or not (root / "latest.json").exists()) and self.preview_root:
            root = self.preview_root
            stage = "internal_trial"
        if root is None or not root.is_dir():
            raise PurchaseSnapshotError("snapshot_not_configured", unavailable=True)
        if root.is_symlink() or (root / "snapshots").is_symlink():
            raise PurchaseSnapshotError("snapshot_path_symlink")
        pointer_path = root / "latest.json"
        if not pointer_path.exists():
            raise PurchaseSnapshotError("snapshot_not_published", unavailable=True)
        pointer = _json_file(pointer_path)
        snapshot_id = pointer.get("snapshot_id")
        if not isinstance(snapshot_id, str) or not SNAPSHOT_ID.fullmatch(snapshot_id):
            raise PurchaseSnapshotError("snapshot_id_invalid")
        expected_manifest = f"snapshots/{snapshot_id}/manifest.json"
        if pointer.get("manifest") != expected_manifest:
            raise PurchaseSnapshotError("snapshot_pointer_invalid")
        directory = root / "snapshots" / snapshot_id
        if directory.is_symlink():
            raise PurchaseSnapshotError("snapshot_path_symlink")
        manifest = _json_file(directory / "manifest.json")
        if (manifest.get("snapshot_id") != snapshot_id or manifest.get("status") != "validated" or
                manifest.get("source") != "beefun_browser_export" or
                manifest.get("units_verified") is not False or
                manifest.get("coverage_complete") is not False):
            raise PurchaseSnapshotError("manifest_identity_invalid")
        account_scope = manifest.get("account_scope")
        scope_basis = manifest.get("account_scope_basis")
        if account_scope is None and scope_basis is None:
            # Pre-scope snapshots remain readable but cannot inherit or create
            # supplier promises because their ERP tenant was never recorded.
            account_scope = ""
        elif (not isinstance(account_scope, str) or not SHA256.fullmatch(account_scope) or
              scope_basis != "login_company_and_verified_erp_host_v1"):
            raise PurchaseSnapshotError("account_scope_invalid")
        for count_key in ("bytes", "row_count", "web_row_count", "order_count",
                          "supplier_count", "duplicate_candidate_key_groups", "header_row"):
            value = manifest.get(count_key)
            if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                raise PurchaseSnapshotError("manifest_count_invalid")
        if manifest["row_count"] < 1 or manifest["row_count"] > MAX_ROWS:
            raise PurchaseSnapshotError("manifest_count_invalid")
        try:
            captured = datetime.fromisoformat(str(manifest["captured_at_utc"]))
            if captured.tzinfo is None or captured.utcoffset().total_seconds() != 0:
                raise ValueError()
        except (KeyError, TypeError, ValueError):
            raise PurchaseSnapshotError("captured_at_invalid") from None
        if (captured - datetime.now(timezone.utc)).total_seconds() > 300:
            raise PurchaseSnapshotError("captured_at_in_future")
        digest = manifest.get("sha256")
        if not isinstance(digest, str) or not SHA256.fullmatch(digest):
            raise PurchaseSnapshotError("manifest_hash_invalid")
        rows = self._validated_rows(directory / "purchase-unfinished.xlsx", manifest, account_scope)
        age_hours = max(0, (datetime.now(timezone.utc) - captured).total_seconds() / 3600)
        meta = {key: manifest[key] for key in (
            "snapshot_id", "captured_at_utc", "query_start", "query_end", "filter",
            "row_count", "web_row_count", "order_count", "supplier_count",
            "duplicate_candidate_key_groups", "minimum_unfinished", "unfinished_sum",
            "sha256", "bytes",
            "units_verified", "coverage_complete")}
        meta.update(source="beefun_browser_export", age_hours=round(age_hours, 2),
                    stale=age_hours > 18,
                    account_scope=account_scope,
                    account_scope_basis=scope_basis or "legacy_unscoped_review",
                    account_scope_verified=bool(account_scope),
                    publication_stage=stage,
                    status="ready" if account_scope and stage == "formal" else "review",
                    review_reason=("internal_trial_only" if stage == "internal_trial" else
                                   "" if account_scope else "account_scope_missing"))
        return meta, rows

    def status(self, user):
        self._authorize(user)
        try:
            meta, _ = self._load()
        except PurchaseSnapshotError as error:
            meta = {"status": "unavailable" if error.unavailable else "invalid",
                    "reason": error.code, "units_verified": False,
                    "coverage_complete": False}
        # Run status is independent of the last validated business snapshot.
        # Never forward arbitrary exception text, URLs or authentication data.
        collector_root = self.root
        if self.preview_root and (self.root is None or not (self.root / "latest.json").exists()):
            collector_root = self.preview_root
        path = collector_root / "collector-status.json" if collector_root else None
        if path is not None and path.exists():
            try:
                state = _json_file(path)
                if state.get('status') not in {'running','needs_human','failed','succeeded'}:
                    raise PurchaseSnapshotError('collector_status_invalid')
                public = {'status':state['status']}
                for field in ('started_at_utc','updated_at_utc'):
                    value = state.get(field)
                    if value:
                        parsed = datetime.fromisoformat(str(value))
                        if parsed.tzinfo is None:
                            raise ValueError('missing timezone')
                        public[field] = parsed.isoformat()
                reason = state.get('reason','')
                if isinstance(reason,str) and re.fullmatch(r'[a-z0-9_]{0,100}',reason):
                    public['reason'] = reason
                meta['collector'] = public
            except (PurchaseSnapshotError,ValueError,TypeError):
                meta['collector'] = {'status':'failed','reason':'collector_status_invalid'}
        return meta

    def list(self, user, supplier_code="", order_no="", keyword="", page=1):
        self._authorize(user)
        meta, rows = self._load()
        if not isinstance(page, int) or page < 1:
            raise ValueError("页码必须为正整数")
        supplier_code, order_no, keyword = _text(supplier_code), _text(order_no), _text(keyword).casefold()
        rows = [row for row in rows if
                (not supplier_code or row["supplier_code"] == supplier_code) and
                (not order_no or row["order_no"] == order_no) and
                (not keyword or keyword in " ".join(row[name] for name in
                 ("supplier_name", "order_no", "product_code", "product_name")).casefold())]
        start = (page - 1) * 50
        items = rows[start:start + 50]
        if items and meta["account_scope_verified"]:
            with self.store.connect() as db:
                keys = [row["line_key"] for row in items]
                promises = {row["line_key"]: row for row in db.execute(
                    "SELECT line_key,promised_date,note,version,source_snapshot_id FROM purchase_supplier_promises WHERE account_scope=? AND line_key IN (%s)" %
                    ",".join("?" for _ in keys), [meta["account_scope"], *keys])}
            enriched = []
            for row in items:
                saved = promises.get(row["line_key"])
                # A report's natural key is not an ERP line ID. Even a key that
                # is unique today may now represent the other half of an earlier
                # duplicate pair. Never carry a reply across snapshot boundaries.
                applies = bool(saved and not row["identity_ambiguous"] and
                               saved["source_snapshot_id"] == meta["snapshot_id"])
                needs_confirmation = bool(saved and not applies)
                enriched.append({**row,
                    "promised_date": saved["promised_date"] if applies else "",
                    "promise_note": saved["note"] if applies else "",
                    "promise_version": saved["version"] if saved else 0,
                    "prior_promised_date": saved["promised_date"] if needs_confirmation else "",
                    "prior_promise_note": saved["note"] if needs_confirmation else "",
                    "promise_requires_confirmation": needs_confirmation,
                    "source_snapshot_id": saved["source_snapshot_id"] if saved else ""})
            items = enriched
        elif items:
            items = [{**row, "promised_date": "", "promise_note": "",
                      "promise_version": 0, "prior_promised_date": "", "prior_promise_note": "",
                      "promise_requires_confirmation": False, "source_snapshot_id": ""} for row in items]
        return {**meta, "items": items, "total": len(rows),
                "page": page, "page_size": 50, "has_more": start + 50 < len(rows)}

    def suppliers(self, user):
        self._authorize(user)
        meta, rows = self._load()
        grouped = {}
        for row in rows:
            item = grouped.setdefault(row["supplier_code"],
                {"supplier_code": row["supplier_code"], "supplier_name": row["supplier_name"],
                 "row_count": 0, "order_nos": set(), "ambiguous_rows": 0})
            item["row_count"] += 1
            item["order_nos"].add(row["order_no"])
            item["ambiguous_rows"] += int(row["identity_ambiguous"])
        return {**meta, "items": [{**{key: value for key, value in item.items() if key != "order_nos"},
                                    "order_count": len(item["order_nos"])}
                                 for item in sorted(grouped.values(), key=lambda value: value["supplier_code"])]}

    def reminder_preview(self, user, supplier_code):
        self._authorize(user)
        meta, rows = self._load()
        selected = [row for row in rows if row["supplier_code"] == supplier_code]
        if not selected:
            raise KeyError("供应商不在当前未交快照中")
        # No amounts, customer data, raw remark, or send action crosses this API.
        public = [{key: row[key] for key in ("line_key", "order_no", "product_code",
            "product_name", "unfinished_quantity", "unit", "arrival_date",
            "identity_ambiguous")} for row in selected]
        return {"snapshot_id": meta["snapshot_id"],
                "captured_at_utc": meta["captured_at_utc"],
                "supplier_code": supplier_code, "supplier_name": selected[0]["supplier_name"],
                "rows": public, "draft_only": True, "send_allowed": False,
                "reason": "采购计量单位和历史覆盖尚未核实",
                "units_verified": False, "coverage_complete": False}

    def promises(self, user, supplier_code=""):
        self._authorize(user)
        meta, _ = self._load()
        if not meta["account_scope_verified"]:
            return {"items": [], "account_scope_verified": False,
                    "review_reason": "account_scope_missing"}
        with self.store.connect() as db:
            rows = db.execute("""SELECT * FROM purchase_supplier_promises
                WHERE account_scope=? AND (?='' OR supplier_code=?)
                ORDER BY updated_at DESC,line_key""",
                (meta["account_scope"], _text(supplier_code), _text(supplier_code))).fetchall()
            return {"items": [dict(row) for row in rows],
                    "account_scope_verified": True}

    def promise(self, user, line_key, body):
        self._authorize(user)
        if not isinstance(body, dict):
            raise ValueError("请求必须为 JSON 对象")
        key = _text(body.get("idempotency_key"))
        note = _text(body.get("note"))
        promised = _date(body.get("promised_date"), "promised_date_invalid").isoformat()
        if not key or len(key) > 128 or not note or len(note) > 1000:
            raise ValueError("请填写交期来源说明和操作标识")
        try:
            expected = int(body.get("version"))
        except (TypeError, ValueError):
            raise ValueError("请提供当前交期记录版本") from None
        with self.store.connect(write=True) as db:
            hold = db.execute("SELECT value FROM settings WHERE key='restoration_hold'").fetchone()
            if hold and hold[0] in ('true', '1'):
                raise DailyWorkConflict("备份恢复后需先核对工作记录", "restoration_hold")
            meta, rows = self._load()
            if meta["publication_stage"] == "internal_trial":
                raise DailyWorkConflict("采购快照仍在内部试用，交期回复暂不写入", "purchase_preview_read_only")
            if not meta["account_scope_verified"]:
                raise DailyWorkConflict("采购快照缺少账套身份，请重新采集后记录交期", "purchase_account_scope_review")
            replay = db.execute("""SELECT e.promised_date,e.note,e.source_snapshot_id
                FROM purchase_supplier_promise_events e
                JOIN purchase_supplier_promises p ON p.line_key=e.line_key
                WHERE e.line_key=? AND e.idempotency_key=? AND p.account_scope=?""",
                (line_key, key, meta["account_scope"])).fetchone()
            if replay:
                if replay["source_snapshot_id"] != body.get("snapshot_id"):
                    raise DailyWorkConflict("该操作标识属于旧采购快照，重新确认需使用新的操作标识", "idempotency_conflict")
                if replay["promised_date"] != promised or replay["note"] != note:
                    raise DailyWorkConflict("操作标识已用于其他交期", "idempotency_conflict")
                row = db.execute("SELECT * FROM purchase_supplier_promises WHERE line_key=?", (line_key,)).fetchone()
                return {"replayed": True, "item": dict(row) if row else None}
            if body.get("snapshot_id") != meta["snapshot_id"]:
                raise DailyWorkConflict("采购快照已更新，请刷新后记录交期", "snapshot_version_changed")
            matches = [row for row in rows if row["line_key"] == line_key]
            if len(matches) != 1 or matches[0]["identity_ambiguous"]:
                raise DailyWorkConflict("采购明细身份不唯一，需人工核对", "purchase_line_conflict")
            line = matches[0]
            previous = db.execute("SELECT * FROM purchase_supplier_promises WHERE line_key=? AND account_scope=?",
                                  (line_key, meta["account_scope"])).fetchone()
            if expected != (previous["version"] if previous else 0):
                raise DailyWorkConflict("供应商回复已变化，请刷新", "task_version_changed")
            now = timestamp()
            if previous:
                db.execute("""UPDATE purchase_supplier_promises SET promised_date=?,note=?,
                    version=version+1,updated_by=?,source_snapshot_id=?,updated_at=? WHERE line_key=?""",
                    (promised, note, user["id"], meta["snapshot_id"], now, line_key))
            else:
                db.execute("""INSERT INTO purchase_supplier_promises(
                    line_key,supplier_code,order_no,product_code,warehouse,order_date,
                    promised_date,note,version,updated_by,source_snapshot_id,
                    created_at,updated_at,account_scope)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (line_key, line["supplier_code"], line["order_no"], line["product_code"],
                     line["warehouse"], line["order_date"], promised, note, 1, user["id"],
                     meta["snapshot_id"], now, now, meta["account_scope"]))
            db.execute("""INSERT INTO purchase_supplier_promise_events VALUES (?,?,?,?,?,?,?,?)""",
                (uuid.uuid4().hex, line_key, key, user["id"], promised, note,
                 meta["snapshot_id"], now))
            updated = db.execute("SELECT * FROM purchase_supplier_promises WHERE line_key=?", (line_key,)).fetchone()
            return {"replayed": False, "item": dict(updated)}

    def promise_events(self, user, line_key, page=1):
        self._authorize(user)
        if not isinstance(page, int) or page < 1:
            raise ValueError("页码必须为正整数")
        meta, _ = self._load()
        if not meta["account_scope_verified"]:
            raise DailyWorkConflict("采购快照缺少账套身份，请重新采集", "purchase_account_scope_review")
        with self.store.connect() as db:
            row = db.execute("SELECT 1 FROM purchase_supplier_promises WHERE line_key=? AND account_scope=?",
                             (line_key, meta["account_scope"])).fetchone()
            if not row:
                raise KeyError("供应商交期记录不存在")
            start = (page - 1) * 100
            events = db.execute("""SELECT * FROM purchase_supplier_promise_events
                WHERE line_key=? ORDER BY created_at DESC,rowid DESC LIMIT 101 OFFSET ?""",
                (line_key, start)).fetchall()
            return {"items": [dict(event) for event in events[:100]],
                    "page": page, "has_more": len(events) > 100}


def purchase_unfinished_routes(engine: PurchaseUnfinished):
    bp = Blueprint("purchase_unfinished", __name__)

    @bp.errorhandler(PurchaseSnapshotError)
    def snapshot_error(error):
        return jsonify(ok=False, error="snapshot_unavailable" if error.unavailable else
                       "snapshot_invalid", reason=error.code,
                       message="采购未交快照暂不可用；请核对上次成功快照"), 503 if error.unavailable else 422

    @bp.get("/api/v2/purchase-unfinished/status")
    def status():
        return jsonify(ok=True, **engine.status(g.current_user))

    @bp.get("/api/v2/purchase-unfinished")
    def list_rows():
        return jsonify(ok=True, **engine.list(g.current_user,
            request.args.get("supplier_code", ""), request.args.get("order_no", ""),
            request.args.get("keyword", ""), int(request.args.get("page", "1"))))

    @bp.get("/api/v2/purchase-unfinished/suppliers")
    def suppliers():
        return jsonify(ok=True, **engine.suppliers(g.current_user))

    @bp.get("/api/v2/purchase-unfinished/suppliers/<supplier_code>/reminder-preview")
    def reminder_preview(supplier_code):
        return jsonify(ok=True, **engine.reminder_preview(g.current_user, supplier_code))

    @bp.get("/api/v2/purchase-unfinished/promises")
    def promises():
        return jsonify(ok=True, **engine.promises(g.current_user,
            request.args.get("supplier_code", "")))

    @bp.put("/api/v2/purchase-unfinished/promises/<line_key>")
    def promise(line_key):
        return jsonify(ok=True, **engine.promise(g.current_user, line_key,
            request.get_json(silent=True)))

    @bp.get("/api/v2/purchase-unfinished/promises/<line_key>/events")
    def promise_events(line_key):
        return jsonify(ok=True, **engine.promise_events(g.current_user, line_key,
            int(request.args.get("page", "1"))))

    return bp
