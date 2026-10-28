"""Bounded quotation intake and immutable supplier documents. No ERP writes."""
from io import BytesIO
from pathlib import Path
from hashlib import sha256
from decimal import Decimal, ROUND_HALF_UP
from xml.sax.saxutils import escape
import json
import uuid
import zipfile
from .store import dumps,timestamp
from .purchase_workspace_data import decimal,number,natural,text

HEADERS=['商品编码','单价','税价口径','交期','每包支数','最小采购支数','运费','说明']
FIELDS=['sku','unit_price','tax_basis','delivery_date','pack_size','minimum_quantity','shipping','note']

def schema(db):
    db.execute('''CREATE TABLE IF NOT EXISTS purchase_workspace_documents(
        id TEXT PRIMARY KEY,actor_id TEXT NOT NULL,idempotency_key TEXT NOT NULL,
        fingerprint TEXT NOT NULL,kind TEXT NOT NULL,supplier_id TEXT NOT NULL,
        snapshot TEXT NOT NULL,file_sha256 TEXT NOT NULL,created_at TEXT NOT NULL,
        UNIQUE(actor_id,idempotency_key))''')

def quote_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    w=Workbook();s=w.active;s.title='供应商报价';s.append(HEADERS)
    for i,width in enumerate([20,16,16,18,16,18,14,45],1):
        s.column_dimensions[get_column_letter(i)].width=width
        s.cell(1,i).fill=PatternFill('solid',fgColor='B91C24');s.cell(1,i).font=Font(color='FFFFFF',bold=True)
    s.freeze_panes='A2';s.auto_filter.ref='A1:H1'
    guide=w.create_sheet('填写说明');guide.append(['按当前商品编码逐行填写；税价口径仅填含税、未税或待核对。'])
    guide.append(['价格单位为元/支；不要填包价。单元格公式不自动执行，请粘贴为数值。'])
    guide.append(['上传后须核对匹配结果，再确认保存。无法匹配的商品不会自动创建或替换。'])
    guide.column_dimensions['A'].width=100
    b=BytesIO();w.save(b);b.seek(0);return b

def preview_quotes(engine,user,blob,filename,supplier_id):
    from openpyxl import load_workbook
    from .purchase_workspace import authorize,PurchaseConflict
    authorize(user)
    if len(blob)>10*1024*1024 or not filename.lower().endswith('.xlsx'):
        raise ValueError('请上传不超过 10 MB 的 xlsx 报价表')
    with zipfile.ZipFile(BytesIO(blob)) as z:
        if sum(x.file_size for x in z.infolist())>50*1024*1024:raise ValueError('解压后文件过大')
    w=load_workbook(BytesIO(blob),read_only=True,data_only=False)
    try:
        s=w['供应商报价'] if '供应商报价' in w.sheetnames else w.active
        rows=s.iter_rows();header=[text(c.value) for c in next(rows)]
        if header[:8]!=HEADERS:raise ValueError('表头与报价模板不同，请先下载模板；不会猜测价格和数量列')
        result=[];seen=set()
        with engine.store.read_snapshot(),engine.store.connect() as db:
            token,current=engine._source(db)
            if not current:raise PurchaseConflict('采购数据正在更新，请稍后上传')
            supplier=db.execute('SELECT * FROM purchase_workspace_suppliers WHERE id=? AND active=1',(supplier_id,)).fetchone()
            if not supplier:raise ValueError('请选择当前有效供应商')
            for n,cells in enumerate(rows,2):
                if n>101:raise ValueError('每次最多导入 100 行报价，请拆分文件')
                values=[c.value for c in cells[:8]]
                if not any(v is not None for v in values):continue
                value=dict(zip(FIELDS,values));code=text(value.get('sku')).upper();issues=[]
                if any(c.data_type=='f' for c in cells):issues.append('请将公式转成数值')
                matches=db.execute("SELECT i.id,d.version FROM purchase_workspace_index i JOIN purchase_workspace_decisions d ON d.item_id=i.id WHERE upper(i.product_code)=? AND i.profile_id=?",(code,supplier['profile_id'])).fetchall()
                if len(matches)!=1:issues.append('商品未进入采购范围或编码身份不唯一')
                if code in seen:issues.append('本次表格商品重复，请保留最终报价');
                seen.add(code)
                try:
                    if decimal(value.get('unit_price'))<0:raise ValueError()
                except ValueError:issues.append('单价无效')
                if text(value.get('tax_basis')) not in {'含税','未税','待核对'}:issues.append('税价口径无效')
                result.append({'row':n,'values':{k:text(v) for k,v in value.items()},'issues':issues,
                    'id':matches[0]['id'] if len(matches)==1 else '', 'version':matches[0]['version'] if len(matches)==1 else 0})
        ident=uuid.uuid4().hex;payload={'source_token':token,'supplier_id':supplier_id,'rows':result}
        with engine.store.connect(write=True) as db:
            db.execute('INSERT INTO purchase_workspace_imports VALUES (?,?,?,?,?,?,?)',(ident,user['id'],'quotes',sha256(blob).hexdigest(),Path(filename).name,dumps(payload),timestamp()))
        return {'id':ident,**payload}
    finally:w.close()

def apply_quotes(engine,user,body):
    from .purchase_workspace import authorize,PurchaseConflict
    authorize(user)
    with engine.store.connect() as db:
        row=db.execute("SELECT * FROM purchase_workspace_imports WHERE id=? AND actor_id=? AND kind='quotes'",(body.get('import_id'),user['id'])).fetchone()
        if not row:raise ValueError('报价预览不存在')
        payload=json.loads(row['rows_json'])
    selected=body.get('rows')
    if not isinstance(selected,list) or not selected or len(selected)!=len(set(selected)):raise ValueError('请选择有效报价行')
    lookup={r['row']:r for r in payload['rows']}
    rows=[lookup[n] for n in selected if n in lookup]
    if len(rows)!=len(selected) or any(r['issues'] for r in rows):raise ValueError('包含未匹配或有问题的报价，尚未保存')
    # All rows share the engine transaction and optimistic checks.
    return engine.mutate(user,{'action':'quote','source_token':payload['source_token'],
        'items':[{'id':r['id'],'version':r['version'],'quote':r['values']} for r in rows],
        'supplier_id':payload['supplier_id'],'idempotency_key':body.get('idempotency_key')})

def export_document(engine,user,body,images):
    import fcntl
    from .purchase_workspace import PurchaseConflict
    with (engine.store.runtime/'.purchase-pdf.lock').open('a') as lock:
        try:fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:raise PurchaseConflict('另一份采购文件正在生成，请稍后重试','pdf_busy') from None
        return _export_document(engine,user,body,images)

def _export_document(engine,user,body,images):
    from .purchase_workspace import authorize,PurchaseConflict
    authorize(user);kind=body.get('kind');key=text(body.get('idempotency_key'))
    if kind not in {'inquiry','purchase'} or len(key)<8:raise ValueError('导出类型或操作标识无效')
    fingerprint=sha256(dumps(body).encode()).hexdigest();root=engine.store.runtime/'purchase-documents';root.mkdir(mode=0o700,exist_ok=True)
    if sum(p.stat().st_size for p in root.glob('*.pdf'))>512*1024*1024:raise ValueError('采购文件已超过512MB，请管理员归档后再导出；已有文件保留')
    with engine.store.read_snapshot(),engine.store.connect() as db:
        old=db.execute('SELECT * FROM purchase_workspace_documents WHERE actor_id=? AND idempotency_key=?',(user['id'],key)).fetchone()
        if old:
            if old['fingerprint']!=fingerprint:raise PurchaseConflict('原导出标识不能替换内容')
            return {'id':old['id'],'kind':old['kind'],'sha256':old['file_sha256']}
        token,current=engine._source(db)
        if not current or body.get('source_token')!=token:raise PurchaseConflict('库存已更新，请刷新后重新核对导出范围')
        items=body.get('items',[])
        if not isinstance(items,list) or not 1<=len(items)<=50 or len({x.get('id') for x in items})!=len(items):raise ValueError('每份文件选择 1–50 款商品')
        supplier=db.execute('SELECT * FROM purchase_workspace_suppliers WHERE id=? AND active=1',(body.get('supplier_id'),)).fetchone()
        if not supplier:raise ValueError('请选择有效供应商')
        chosen=body.get('allocation_ids')
        if kind=='purchase' or chosen is not None:
            if not isinstance(chosen,list) or not 1<=len(chosen)<=100 or len(set(chosen))!=len(chosen):
                raise ValueError('请明确勾选本次导出的采购安排，最多100行；不会自动带入历史采购')
        lookup={x['id']:x for x in engine._items(db)};lines=[]
        for selected in items:
            x=lookup.get(selected['id'])
            if not x or x['version']!=selected.get('version'):raise PurchaseConflict('商品安排已变化，请重新核对；没有导出')
            allocations=[a for a in x['allocations'] if a['supplier_id']==supplier['id'] and (a['state']=='draft' if kind=='inquiry' else a['state'] in {'draft','ordered','awaiting_audit','erp_counted'})]
            if chosen is not None:allocations=[a for a in allocations if a['id'] in chosen]
            if not allocations:raise ValueError(x['sku']+' 没有该供应商对应的采购安排')
            if kind=='purchase' and (x.get('source_missing') or not x.get('sku_id') or (x['review_reasons'] and not x['review_ack'])):
                raise ValueError(x['sku']+' 请先核对商品身份及需求原因，再导出采购单')
            for a in allocations:
                if kind=='purchase' and a['state']=='draft' and not a['details'].get('quote_id'):
                    raise ValueError(x['sku']+' 请先在供应商报价中采纳报价，再导出采购单')
                lines.append({**{k:x[k] for k in ('id','product_id','sku','product_name','spec','model','material','kind','unit','has_image')},
                    'allocation_id':a['id'],'quantity':a['quantity'],'details':a['details'],'version':x['version'],'state':a['state']})
        if chosen is not None and {x['allocation_id'] for x in lines}!=set(chosen):
            raise PurchaseConflict('所选采购明细不属于本次商品或供应商范围')
        if len(lines)>100:raise ValueError('每份文件最多100行采购明细，请分批导出')
        lines.sort(key=lambda x:(natural(x['product_name']),natural(x['spec']),natural(x['model']),natural(x['sku']),x['allocation_id']))
        snapshot={'kind':kind,'supplier_name':supplier['name'],'supplier_id':supplier['id'],'source_token':token,'created_at':timestamp(),'lines':lines}
    ident=uuid.uuid4().hex;path=root/(ident+'.pdf');attachments=[]
    if kind=='purchase': prepare_order_snapshot(snapshot,body,ident)
    for item in {x['id']:x for x in lines}.values():
        if item['kind']!='非标件':continue
        if not item['has_image']:raise ValueError(item['sku']+' 是非标件但尚无图纸，请补齐图纸后再导出')
        try:
            metadata=images.list_for_product(item['product_id'])['items']
            if metadata:
                for entry in metadata:
                    image=images.content(entry['id'])
                    attachments.append((item,Path(image.path).read_bytes()))
                    if len(attachments)>50 or sum(len(content) for _,content in attachments)>20*1024*1024:
                        raise ValueError('图纸超过50张或20MB，请分批导出')
            else:
                raise ValueError('主图引用未同步')
        except Exception:raise ValueError(item['sku']+' 图纸暂时无法取得，请先打开图纸核对后重试') from None
    snapshot['images']=[{'sku':x['sku'],'sha256':sha256(data).hexdigest()} for x,data in attachments]
    render_pdf(path,snapshot,attachments)
    digest=sha256(path.read_bytes()).hexdigest()
    try:
        with engine.store.connect(write=True) as db:
            fresh,_=engine._source(db)
            if fresh!=token:raise PurchaseConflict('导出期间库存更新，请重新核对范围；文件未发布')
            for target in items:
                r=db.execute('SELECT version FROM purchase_workspace_decisions WHERE item_id=?',(target['id'],)).fetchone()
                if not r or r[0]!=target['version']:raise PurchaseConflict('导出期间采购安排变化，文件未发布')
            db.execute('INSERT INTO purchase_workspace_documents VALUES (?,?,?,?,?,?,?,?,?)',(ident,user['id'],key,fingerprint,kind,supplier['id'],dumps(snapshot),digest,timestamp()))
    except Exception:
        path.unlink(missing_ok=True);raise
    return {'id':ident,'kind':kind,'sha256':digest,'lines':len(lines),'images':len(attachments)}

def prepare_order_snapshot(snapshot,body,ident):
    """Freeze template metadata and tax-inclusive values; never infer a missing price."""
    from datetime import datetime,date
    from zoneinfo import ZoneInfo
    template=Path(__file__).resolve().parents[1]/'assets/templates/purchase-template.json'
    snapshot['buyer']=json.loads(template.read_text())
    snapshot['order_date']=text(body.get('order_date')) or datetime.now(ZoneInfo('Asia/Shanghai')).date().isoformat()
    date.fromisoformat(snapshot['order_date'])
    snapshot['order_number']=text(body.get('order_number'))[:60] or 'CG-'+snapshot['order_date'].replace('-','')+'-'+ident[:8].upper()
    snapshot['supplier_note']=text(body.get('supplier_note'))[:500]
    dates={text(x['details'].get('delivery_date')) for x in snapshot['lines']}
    snapshot['delivery_date']=text(body.get('delivery_date'))[:80] or (next(iter(dates)) if len(dates)==1 and '' not in dates else '详见各行交期' if any(dates) else '交期待确认')
    total=Decimal(0)
    for x in snapshot['lines']:
        d=x['details']
        if not text(d.get('unit_price')) or d.get('tax_basis') not in {'含税','未税'}:raise ValueError(x['sku']+' 单价或税价口径尚未确认')
        price=decimal(d['unit_price']);quantity=decimal(x['quantity'])
        if price<0 or quantity<=0:raise ValueError(x['sku']+' 价格或数量无效')
        factor=Decimal('1.13') if d['tax_basis']=='未税' else Decimal(1)
        # Recorded purchases retain the actual paid amount; draft quotes use current quantity.
        amount=decimal(d['amount']) if x['state']!='draft' else price*quantity
        x['tax_unit_price']=number((price*factor).quantize(Decimal('.000001'),rounding=ROUND_HALF_UP))
        tax_amount=(amount*factor).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        x['tax_amount']=format(tax_amount,'.2f');total+=tax_amount
    snapshot['tax_total']=format(total,'.2f')

def render_pdf(path,snapshot,attachments):
    if snapshot['kind']=='purchase':
        from .purchase_order_pdf import render_order
        return render_order(path,snapshot,attachments)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image,KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor,white
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    if 'ULChinese' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('ULChinese',str(Path(__file__).resolve().parents[1]/'assets/fonts/NotoSansSC.ttf')))
    body=ParagraphStyle('body',fontName='ULChinese',fontSize=9,leading=14,wordWrap='CJK')
    title=ParagraphStyle('title',parent=body,fontSize=22,leading=29,textColor=HexColor('#b91c24'),spaceAfter=18)
    sub=ParagraphStyle('sub',parent=body,fontSize=12,leading=18,spaceAfter=12)
    p=lambda v:Paragraph(escape(text(v)),body)
    is_order=snapshot['kind']=='purchase';story=[Paragraph('UL随行  /  '+('采购单' if is_order else '采购询价单'),title),Paragraph(escape(snapshot['supplier_name']),sub),p('编制时间：'+snapshot['created_at']),Spacer(1,15)]
    header=['序号','商品编码 / 名称','规格 / 型号','数量（支）']+(['单价 / 税价','金额（元）'] if is_order else ['报价 / 交期'])
    rows=[[p(x) for x in header]]
    for i,x in enumerate(snapshot['lines'],1):
        d=x['details'];r=[str(i),x['sku']+'\n'+x['product_name'],x['spec']+'\n'+x['model'],x['quantity']]
        r+=([d.get('unit_price','')+'\n'+d.get('tax_basis',''),d.get('amount','')] if is_order else ['请填写'])
        rows.append([Paragraph(escape(text(v)).replace('\n','<br/>'),body) for v in r])
    widths=[36,154,116,60,82,55] if is_order else [36,180,130,65,92]
    t=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#f5e8e9')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),10),('TOPPADDING',(0,0),(-1,-1),10),('LINEBELOW',(0,0),(-1,-1),.4,HexColor('#e4e7ec'))]));story.append(t)
    story+=[Spacer(1,16),p('本文件按当前采购安排固定生成；客户信息及历史采购价格不对外展示。'+('实际价格以双方确认的税价、运费及折扣为准。' if is_order else '请回复单价、税价口径、包装支数及交期。'))]
    for item,data in attachments:
        story += [PageBreak(),Paragraph('图纸附件 · '+escape(item['sku']),sub),p(item['product_name']+' '+item['spec']+' '+item['model']),p('来源：ERP 商品主图；请核对图号与版本。'),Spacer(1,15)]
        w,h=ImageReader(BytesIO(data)).getSize();factor=min(503/w,610/h)
        story.append(Image(BytesIO(data),width=w*factor,height=h*factor))
    def footer(c,doc):
        c.setFont('ULChinese',8);c.setFillColor(HexColor('#9299a4'));c.drawString(46,27,'江苏优螺工业科技有限公司');c.drawRightString(A4[0]-46,27,str(doc.page))
    SimpleDocTemplate(str(path),pagesize=A4,rightMargin=46,leftMargin=46,topMargin=42,bottomMargin=46).build(story,onFirstPage=footer,onLaterPages=footer)
