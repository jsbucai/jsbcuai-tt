"""Reject incomplete workbooks without replacing the last valid snapshot."""

from datetime import date
from pathlib import Path
import json
import tempfile
import unittest

from openpyxl import Workbook

from purchase_export_collector.collector import publish_snapshot
from purchase_export_collector.validator import (
    ExportValidationError, REQUIRED_COLUMNS, SHEET_NAME, validate_export,
)


START = date(1900, 1, 1)
END = date(2026, 10, 6)


def write_report(path: Path, *, total=True, total_quantity=10) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.cell(1, 1, SHEET_NAME)
    sheet.cell(2, 1, "开始时间")
    sheet.cell(2, 2, START.isoformat())
    sheet.cell(3, 1, "结束时间")
    sheet.cell(3, 2, END.isoformat())
    values = {
        "行号": 1, "单位编号": "SUPPLIER-1", "单位名称": "测试供应商",
        "商品编号": "PRODUCT-1", "商品名称": "测试商品", "订单编号": "PO-1",
        "仓库": "测试仓库", "订货日期": "2026-10-01", "到货日期": "2026-10-10",
        "订货数量": 10, "完成数量": 0, "终止完成数量": 0, "未完成数量": 10,
    }
    for index, name in enumerate(REQUIRED_COLUMNS, 1):
        sheet.cell(5, index, name)
        sheet.cell(6, index, values[name])
    if total:
        sheet.cell(7, REQUIRED_COLUMNS.index("行号") + 1, "合计")
        sheet.cell(7, REQUIRED_COLUMNS.index("未完成数量") + 1, total_quantity)
    book.save(path)
    book.close()


class ExportIntegrityTests(unittest.TestCase):
    def test_missing_footer_is_rejected_even_when_all_detail_counts_match(self):
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / "incomplete.xlsx"
            write_report(report, total=False)
            with self.assertRaisesRegex(ExportValidationError, "missing_total_row"):
                validate_export(report, start=START, end=END, expected_rows=1)

    def test_complete_footer_is_required_and_its_total_must_match(self):
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / "report.xlsx"
            write_report(report)
            verified = validate_export(report, start=START, end=END, expected_rows=1)
            self.assertEqual(verified["row_count"], 1)
            self.assertEqual(verified["unfinished_sum"], "10")
            write_report(report, total_quantity=9)
            with self.assertRaisesRegex(ExportValidationError, "total_unfinished_mismatch"):
                validate_export(report, start=START, end=END, expected_rows=1)

    def test_incomplete_staged_file_cannot_replace_last_success(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            good, incomplete = root / "complete.xlsx", root / "incomplete.xlsx"
            write_report(good)
            write_report(incomplete, total=False)
            verified = validate_export(good, start=START, end=END, expected_rows=1)
            output = root / "isolated-snapshots"
            manifest = publish_snapshot(good, output, verified,
                                        web_rows=1, account_scope="a" * 64)
            previous_pointer = (output / "latest.json").read_bytes()
            previous_manifest = manifest.read_bytes()
            previous_workbook = (manifest.parent / "purchase-unfinished.xlsx").read_bytes()
            with self.assertRaisesRegex(ExportValidationError, "missing_total_row"):
                publish_snapshot(incomplete, output, verified,
                                 web_rows=1, account_scope="a" * 64)
            self.assertEqual((output / "latest.json").read_bytes(), previous_pointer)
            self.assertEqual(manifest.read_bytes(), previous_manifest)
            self.assertEqual((manifest.parent / "purchase-unfinished.xlsx").read_bytes(),
                             previous_workbook)
            self.assertEqual(json.loads(previous_pointer)["snapshot_id"], manifest.parent.name)
            self.assertFalse(list((output / "snapshots").glob(".staging-*")))
            self.assertEqual(len(list((output / "snapshots").iterdir())), 1)


if __name__ == "__main__":
    unittest.main()
