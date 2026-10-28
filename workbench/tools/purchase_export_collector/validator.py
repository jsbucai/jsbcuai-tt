"""Validate the exact unfinished-purchase export before publishing a snapshot."""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
import warnings

from openpyxl import load_workbook


SHEET_NAME = "进货订单统计列表"
REQUIRED_COLUMNS = (
    "行号", "单位编号", "单位名称", "商品编号", "商品名称", "订单编号",
    "仓库", "订货日期", "到货日期", "订货数量", "完成数量",
    "终止完成数量", "未完成数量",
)
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000


class ExportValidationError(ValueError):
    """The file cannot safely replace the last successful snapshot."""


def _date(value: object, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            try:
                return date.fromisoformat(value)
            except ValueError:
                pass
    raise ExportValidationError(f"invalid_{label}")


def _quantity(value: object, label: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ExportValidationError(f"invalid_{label}")
    try:
        number = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError):
        raise ExportValidationError(f"invalid_{label}") from None
    if not number.is_finite():
        raise ExportValidationError(f"invalid_{label}")
    return number


def validate_export(
    path: Path,
    *,
    start: date,
    end: date,
    min_unfinished: Decimal = Decimal("0.1"),
    expected_rows: int | None = None,
) -> dict[str, object]:
    """Read the original workbook without changing its rows or unit values.

    The vendor workbook has an incorrect OOXML worksheet dimension (A1), so
    openpyxl's read-only mode silently hides all but column A. Use normal
    loading, with a file-size limit, and never save this workbook.
    """
    if start > end or min_unfinished <= 0:
        raise ExportValidationError("invalid_query_scope")
    if not path.is_file() or path.stat().st_size == 0:
        raise ExportValidationError("empty_or_missing_file")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ExportValidationError("file_too_large")
    if expected_rows is not None and expected_rows < 0:
        raise ExportValidationError("invalid_web_count")
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style")
            workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception:
        raise ExportValidationError("invalid_xlsx") from None
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise ExportValidationError("wrong_sheet")
        sheet = workbook.active
        if sheet.max_row > MAX_ROWS + 20:
            raise ExportValidationError("too_many_rows")
        if sheet.cell(1, 1).value != SHEET_NAME:
            raise ExportValidationError("wrong_report_title")
        if sheet.cell(2, 1).value != "开始时间" or sheet.cell(3, 1).value != "结束时间":
            raise ExportValidationError("missing_query_dates")
        if _date(sheet.cell(2, 2).value, "start_date") != start:
            raise ExportValidationError("start_date_mismatch")
        if _date(sheet.cell(3, 2).value, "end_date") != end:
            raise ExportValidationError("end_date_mismatch")
        header_row = None
        for row_num in range(4, min(sheet.max_row, 15) + 1):
            headers = [sheet.cell(row_num, col).value for col in range(1, sheet.max_column + 1)]
            if "订单编号" in headers and "未完成数量" in headers:
                header_row = row_num
                break
        if header_row is None:
            raise ExportValidationError("missing_detail_header")
        header = [sheet.cell(header_row, col).value for col in range(1, sheet.max_column + 1)]
        if any(not isinstance(name, str) or not name.strip() for name in header):
            raise ExportValidationError("blank_header")
        if len(header) != len(set(header)) or not set(REQUIRED_COLUMNS).issubset(header):
            raise ExportValidationError("missing_or_duplicate_columns")
        col = {name: header.index(name) + 1 for name in header}
        orders: set[str] = set()
        suppliers: set[str] = set()
        keys: Counter[tuple[str, str, str, str]] = Counter()
        row_count = 0
        row_numbers: set[int] = set()
        min_quantity: Decimal | None = None
        unfinished_sum = Decimal("0")
        total_row_seen = False
        for row_num in range(header_row + 1, sheet.max_row + 1):
            values = {name: sheet.cell(row_num, index).value for name, index in col.items()}
            if all(value is None or value == "" for value in values.values()):
                continue
            serial = values["行号"]
            if serial == "合计" and row_num == sheet.max_row and not total_row_seen:
                total_row_seen = True
                if _quantity(values["未完成数量"], "total_unfinished_quantity") != unfinished_sum:
                    raise ExportValidationError("total_unfinished_mismatch")
                continue
            if isinstance(serial, bool) or not isinstance(serial, (int, float)) or int(serial) != serial or serial < 1:
                raise ExportValidationError("unexpected_non_detail_row")
            serial = int(serial)
            if serial in row_numbers:
                raise ExportValidationError("duplicate_row_number")
            row_numbers.add(serial)
            for name in ("单位编号", "单位名称", "商品编号", "商品名称", "订单编号", "仓库"):
                if not str(values[name] or "").strip():
                    raise ExportValidationError("blank_identity_field")
            order_date = _date(values["订货日期"], "order_date")
            if not start <= order_date <= end:
                raise ExportValidationError("order_date_outside_query")
            if values["到货日期"] not in (None, ""):
                _date(values["到货日期"], "arrival_date")
            unfinished = _quantity(values["未完成数量"], "unfinished_quantity")
            if unfinished < min_unfinished:
                raise ExportValidationError("unfinished_below_filter")
            for name in ("订货数量", "完成数量", "终止完成数量"):
                if _quantity(values[name], name) < 0:
                    raise ExportValidationError("negative_quantity")
            min_quantity = unfinished if min_quantity is None else min(min_quantity, unfinished)
            unfinished_sum += unfinished
            orders.add(str(values["订单编号"]).strip())
            suppliers.add(str(values["单位编号"]).strip())
            keys[(str(values["单位编号"]).strip(), str(values["订单编号"]).strip(),
                  str(values["商品编号"]).strip(), str(values["仓库"]).strip())] += 1
            row_count += 1
        if row_count == 0:
            raise ExportValidationError("zero_rows_requires_manual_review")
        if not total_row_seen:
            raise ExportValidationError("missing_total_row")
        if row_numbers != set(range(1, row_count + 1)):
            raise ExportValidationError("row_numbers_not_contiguous")
        if expected_rows is not None and row_count != expected_rows:
            raise ExportValidationError("web_file_count_mismatch")
        return {
            "sheet": SHEET_NAME,
            "header_row": header_row,
            "columns": header,
            "query_start": start.isoformat(),
            "query_end": end.isoformat(),
            "filter": f"未完成数量 >= {min_unfinished}",
            "row_count": row_count,
            "order_count": len(orders),
            "supplier_count": len(suppliers),
            "duplicate_candidate_key_groups": sum(count > 1 for count in keys.values()),
            "minimum_unfinished": str(min_quantity),
            "unfinished_sum": str(unfinished_sum),
            "sha256": sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
            "units_verified": False,
            "coverage_complete": False,
        }
    finally:
        workbook.close()
